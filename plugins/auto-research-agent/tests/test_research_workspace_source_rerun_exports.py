"""Synthetic regressions for the saved-source rerun projection."""

import io
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parents[1] / "cli"))

import test_research_workspace_view as view_fixture
import test_research_workspace_closeout as closeout_fixture
from research_workspace.projection import validate_index
from research_workspace.source_rerun import attach_rerun, rerun_files
from research_workspace.view import render_view, write_workspace
from research_workspace.wiki import wiki_files
from stage1_deliverable.common import canonical, inventory, sha
from stage1_deliverable.views import bibtex


RAW = {
    "source-0": b"<html><title>Evidence</title><body>saved full text</body></html>",
    "source-1": b"<html><title>literal title</title><body>saved abstract</body></html>",
    "source-2": b"<html><title>Evidence</title><body>second saved source</body></html>",
}
TEXT = {
    "source-0": b"Evidence saved full text",
    "source-1": b"literal title saved abstract",
    "source-2": b"Evidence second saved source",
}


def _receipt(source_id, status, evidence_level):
    raw_hash = sha(RAW[source_id])
    return {
        "status": status,
        "evidence_level": evidence_level,
        "attempts": [
            {
                "sequence": 1,
                "purpose": "synthetic offline fixture",
                "url": f"https://example.test/{source_id}",
                "final_url": f"https://example.test/{source_id}",
                "http_status": 200,
                "outcome": status,
                "response_truncated": False,
                "raw_path": "saved-response.html",
                "raw_sha256": raw_hash,
                "error": None,
            }
        ],
    }


def fixture_index():
    index = view_fixture.fixture_index()
    index["claims"] = [
        {
            "claim_id": "historical-claim",
            "work_id": "fixture-0",
            "version_id": "version-0",
            "text": "A historical synthetic claim remains unmodified.",
            "support": "Unknown",
        }
    ]
    index["coverage"] = [
        {
            "need_id": "synthetic-need",
            "status": "unresolved",
            "coverage": "Unknown",
        }
    ]
    index["sources"] = []
    for number, (status, level) in enumerate(
        (("parse-error", "metadata"), ("abstract-only", "abstract"))
    ):
        source_id = f"source-{number}"
        receipt = _receipt(source_id, status, level)
        index["sources"].append(
            {
                "source_id": source_id,
                "work_id": f"fixture-{number}",
                "version_id": f"version-{number}",
                "result_sha256": sha(canonical(receipt)),
                "receipt": receipt,
            }
        )
    validate_index(index)
    return index


def refresh_bibliography(index):
    index["bibliography"]["entries"] = [
        {
            "work_id": paper["work_id"],
            "version_id": paper["version_id"],
            "bibtex": bibtex({"papers": [paper]}).decode("utf-8"),
        }
        for paper in index["papers"]
    ]
    index["bibliography"]["all_bibtex"] = bibtex(index).decode("utf-8")
    validate_index(index)


def bind_observed_metadata(row, **values):
    row["source_metadata"] = {"title": "Evidence", **values}
    for key, value in values.items():
        row["metadata"][key] = value
        row["metadata_provenance"][key] = {
            "status": "observed-in-identity-matched-saved-source",
            "value": value,
            "locators": row["reading"]["locators"],
        }


def _row(index, number):
    source = index["sources"][number]
    source_id = source["source_id"]
    raw_hash = sha(RAW[source_id])
    text_hash = sha(TEXT[source_id])
    level = "full-text" if number == 0 else "abstract"
    parser_hash = "f" * 64
    base_hash = sha(canonical(index))
    row = {
        "work_id": source["work_id"],
        "version_id": source["version_id"],
        "source_id": source_id,
        "original_result_sha256": source["result_sha256"],
        "previous_status": source["receipt"]["status"],
        "previous_evidence_level": source["receipt"]["evidence_level"],
        "original_attempts": deepcopy(source["receipt"]["attempts"]),
        "raw_path": f"sources/{source_id}/raw.html",
        "raw_sha256": raw_hash,
        "source_version": "sha256:" + raw_hash,
        "final_url": f"https://example.test/{source_id}",
        "selected_sequence": 1,
        "extracted_path": f"sources/{source_id}/extracted.txt",
        "reading": {
            "status": "extracted",
            "evidence_level": level,
            "identity_status": "consistent",
            "characters": len(TEXT[source_id].decode("utf-8")),
            "text_sha256": text_hash,
            "locators": [
                {
                    "start": 0,
                    "end": len(TEXT[source_id].decode("utf-8")) if number else 8,
                    "type": "text",
                    "value": "abstract" if number else "body",
                }
            ],
            "diagnostics": {},
            "error": None,
        },
        "metadata": {
            "journal": "Unknown",
            "volume": None,
            "issue": None,
            "pages": None,
            "doi": None,
        },
        "metadata_provenance": {
            "journal": {"status": "original-verbatim", "value": "Unknown"}
        },
    }
    row["attempt_id"] = sha(
        canonical(
            {
                "base": base_hash,
                "parser": parser_hash,
                "work": row["work_id"],
                "version": row["version_id"],
                "raw": raw_hash,
            }
        )
    )
    return row


def fixture_manifest(index):
    rows = [_row(index, 0), _row(index, 1)]
    files = {}
    for row in rows:
        source_id = row["source_id"]
        for name, raw in (
            (row["raw_path"], RAW[source_id]),
            (row["extracted_path"], TEXT[source_id]),
        ):
            files[name] = {"sha256": sha(raw), "bytes": len(raw)}
    return {
        "data": {
            "kind": "Stage1SavedSourceRerun",
            "schema_version": "1.0.0",
            "base_index_sha256": sha(canonical(index)),
            "parser_runtime": {
                "fixture": "offline-synthetic",
                "parser_sha256": "f" * 64,
            },
            "rows": rows,
            "research_execution": False,
            "scientific_judgments_changed": False,
            "official_stage2_import_eligible": False,
            "new_searches": 0,
            "new_downloads": 0,
            "quality_score": None,
        },
        "files": dict(sorted(files.items())),
    }


def write_manifest(root, manifest):
    for name in manifest["files"]:
        source_id = Path(name).parts[1]
        raw = RAW[source_id] if name.endswith("raw.html") else TEXT[source_id]
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(raw)
    raw = canonical(manifest)
    (root / "source-rerun-manifest.json").write_bytes(raw)
    return sha(raw)


class SavedSourceRerunTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="source-rerun-test-", dir=Path(tempfile.gettempdir()).resolve()
        )
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.index = fixture_index()

    def attach(self, manifest=None):
        manifest = manifest or fixture_manifest(self.index)
        expected = write_manifest(self.root, manifest)
        return attach_rerun(self.index, self.root, expected)

    def test_cross_format_ids_and_rebuild_are_deterministic(self):
        manifest = fixture_manifest(self.index)
        expected = write_manifest(self.root, manifest)
        first = attach_rerun(self.index, self.root, expected)
        second = attach_rerun(self.index, self.root, expected)
        self.assertEqual(canonical(first), canonical(second))
        first_files = rerun_files(first, self.root)
        second_files = rerun_files(second, self.root)
        self.assertEqual(first_files, second_files)

        workbook = load_workbook(
            io.BytesIO(first_files["source-rerun/catalog.xlsx"]), read_only=True
        )
        rows = list(workbook["SourceRerun"].iter_rows(values_only=True))
        headers = list(rows[0])
        identities = {
            (
                row[headers.index("work_id")],
                row[headers.index("version_id")],
                row[headers.index("source_id")],
            )
            for row in rows[1:]
        }
        expected_ids = {
            ("fixture-0", "version-0", "source-0"),
            ("fixture-1", "version-1", "source-1"),
        }
        self.assertEqual(identities, expected_ids)
        bibtex = first_files["source-rerun/references.bib"].decode("utf-8")
        markdown = first_files["source-rerun/report.md"].decode("utf-8")
        for work_id, version_id, source_id in expected_ids:
            self.assertIn("@misc{" + work_id, bibtex)
            self.assertIn("Version " + version_id, bibtex)
            self.assertIn(work_id, markdown)
            self.assertIn(version_id, markdown)
            self.assertIn(source_id, markdown)

    def test_mixed_read_success_and_failure_export_without_missing_columns(self):
        manifest = fixture_manifest(self.index)
        row = manifest["data"]["rows"][1]
        manifest["files"].pop(row.pop("extracted_path"))
        row["reading"].update(
            status="failed-engineering",
            characters=0,
            text_sha256=None,
            locators=[],
            evidence_level="metadata",
            error={"type": "ValueError", "message": "synthetic parse failure"},
        )
        row["reading"].pop("observed_identity", None)
        manifest["data"]["rows"][0]["reading"]["observed_identity"] = {
            "title": "Evidence",
            "doi": "",
        }
        result = self.attach(manifest)
        exported = rerun_files(result, self.root)
        self.assertIn(b"failed-engineering", exported["source-rerun/catalog.csv"])
        self.assertIn(b"synthetic parse failure", exported["source-rerun/report.md"])

    def test_conflicting_observed_doi_is_retained_but_not_promoted(self):
        self.index["papers"][0]["doi"] = "10.1000/original"
        refresh_bibliography(self.index)
        manifest = fixture_manifest(self.index)
        row = manifest["data"]["rows"][0]
        row["reading"].update(
            status="identity-mismatch",
            identity_status="mismatch",
            observed_identity={"title": "Evidence", "doi": "10.1000/conflict"},
        )
        row["source_metadata"] = {
            "title": "Evidence",
            "doi": "10.1000/conflict",
            "journal": "Wrong Journal",
        }
        row["metadata"]["doi"] = "10.1000/original"
        projected = self.attach(manifest)
        accepted = projected["source_rerun"]["data"]["rows"][0]
        self.assertEqual(
            accepted["reading"]["observed_identity"]["doi"], "10.1000/conflict"
        )
        self.assertEqual(accepted["metadata"]["doi"], "10.1000/original")
        self.assertEqual(accepted["metadata"]["journal"], "Unknown")
        generated = projected["source_rerun"]["bibliography"]["entries"][0]["bibtex"]
        self.assertEqual(generated.count("  doi = {"), 1)
        self.assertNotIn("10.1000/conflict", generated)

    def test_equivalent_doi_spellings_emit_one_bibliography_field(self):
        self.index["papers"][0]["doi"] = "doi:10.1000/ABC"
        refresh_bibliography(self.index)
        manifest = fixture_manifest(self.index)
        row = manifest["data"]["rows"][0]
        bind_observed_metadata(row, doi="https://doi.org/10.1000/abc")
        projected = self.attach(manifest)
        accepted = projected["source_rerun"]["data"]["rows"][0]
        self.assertEqual(accepted["metadata"]["doi"], "https://doi.org/10.1000/abc")
        generated = projected["source_rerun"]["bibliography"]["entries"][0]["bibtex"]
        self.assertEqual(generated.count("  doi = {"), 1)
        self.assertIn("doi:10.1000/ABC", generated)

    def test_v1_to_v3_write_and_render_rebuild_identical_bytes(self):
        rerun_root = self.root / "source-rerun"
        rerun_root.mkdir()
        manifest = fixture_manifest(self.index)
        expected = write_manifest(rerun_root, manifest)
        projected = attach_rerun(self.index, rerun_root, expected)
        index_path = self.root / "workspace-index.json"
        index_path.write_bytes(canonical(projected))
        reference_root = Path(__file__).parents[1] / "references/research-workspace"
        direct = self.root / "direct-view"
        replay = self.root / "replayed-view"
        write_workspace(
            projected,
            reference_root,
            direct,
            source_rerun_root=rerun_root,
        )
        render_view(index_path, replay, reference_root, sha(index_path.read_bytes()))
        self.assertEqual(inventory(direct), inventory(replay))
        for relative in inventory(direct):
            self.assertEqual(
                (direct / relative).read_bytes(), (replay / relative).read_bytes()
            )

    def test_v2_repair_only_retains_producer_adapter_bindings(self):
        repaired = closeout_fixture.CloseoutTests()
        repaired.setUp()
        self.addCleanup(repaired.doCleanups)
        self.assertEqual(repaired.index["schema_version"], "2.0.0")
        self.assertNotIn("source_rerun", repaired.index)
        reference = Path(__file__).parents[1] / "references/research-workspace"
        output = Path(repaired.repair.temporary.name).resolve() / "v2-view"
        receipt = write_workspace(repaired.index, reference, output)
        producer = Path(__file__).parents[1] / "cli/stage1_deliverable"
        for name in ("views.py", "common.py", "records.py", "sources.py", "package.py"):
            with self.subTest(name=name):
                self.assertEqual(
                    receipt["adapter_sources"]["stage1_deliverable/" + name],
                    sha((producer / name).read_bytes()),
                )

    def test_multiple_sources_use_first_canonical_metadata_and_keep_all_rows(self):
        receipt = _receipt("source-2", "available", "full-text")
        self.index["sources"].append(
            {
                "source_id": "source-2",
                "work_id": "fixture-0",
                "version_id": "version-0",
                "result_sha256": sha(canonical(receipt)),
                "receipt": receipt,
            }
        )
        self.index["papers"][0]["source_ids"].append("source-2")
        validate_index(self.index)
        manifest = fixture_manifest(self.index)
        third = _row(self.index, 2)
        bind_observed_metadata(manifest["data"]["rows"][0], volume="11")
        bind_observed_metadata(third, volume="22")
        manifest["data"]["rows"].append(third)
        for name, raw in (
            (third["raw_path"], RAW["source-2"]),
            (third["extracted_path"], TEXT["source-2"]),
        ):
            manifest["files"][name] = {"sha256": sha(raw), "bytes": len(raw)}
        projected = self.attach(manifest)
        files = rerun_files(projected, self.root)
        bibliography = projected["source_rerun"]["bibliography"]["entries"][0]["bibtex"]
        self.assertIn("  volume = {11}", bibliography)
        self.assertNotIn("  volume = {22}", bibliography)
        self.assertEqual(
            len(files["source-rerun/catalog.csv"].decode("utf-8-sig").splitlines()),
            4,
        )
        note = next(
            value.decode("utf-8")
            for name, value in wiki_files(projected).items()
            if name != "wiki/README.md"
            and '"work_id": "fixture-0"' in value.decode("utf-8")
        )
        rerun_note = note.split("## New saved-source read attempt", 1)[1]
        self.assertIn('"source_id": "source-0"', rerun_note)
        self.assertIn('"volume": "11"', rerun_note)
        self.assertNotIn('"volume": "22"', rerun_note)


if __name__ == "__main__":
    unittest.main()
