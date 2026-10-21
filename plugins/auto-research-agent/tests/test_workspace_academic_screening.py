"""Saved academic decisions remain independent of technical source selection."""

import csv
import io
import json
import sys
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook
from markdown_it import MarkdownIt

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parents[1] / "cli"))

from test_research_workspace_view import fixture_index
from research_workspace.literature_selection import (
    derive_literature_selection,
    selection_files,
)
from stage1_deliverable.common import canonical, sha
from stage1_deliverable.views import bibtex


def decision(number, work=0, status="include", reason="Recorded topic relevance"):
    return dict(
        decision_id=f"e{number}",
        work_id=f"fixture-{work}",
        version_id=f"version-{work}",
        status=status,
        reason=reason,
        observed_at="2026-10-04T00:00:00Z",
    )


def academic(index):
    return json.loads(
        selection_files(index, include_academic_screening=True)[
            "literature/academic_screening.json"
        ]
    )


class AcademicScreeningTests(unittest.TestCase):
    def test_shared_reason_counts_distinct_works_without_endorsing_template(self):
        index = fixture_index()
        index["screening"] = [decision(1), decision(2), decision(3, work=1)]
        result = academic(index)
        self.assertEqual(result["counts"]["distinct_current_reasons"], 1)
        self.assertEqual(
            [r["reason_shared_work_count"] for r in result["rows"]], [2, 2]
        )
        self.assertTrue(all(not r["scientific_assessment"] for r in result["rows"]))

    def test_academic_decisions_survive_without_rerun_or_claim_changes(self):
        index = fixture_index()
        index["screening"] = [decision(9), decision(10, 1, "exclude", "Outside scope")]
        index["claims"] = [{"support": "Unknown"}]
        before, technical = canonical(index), derive_literature_selection(index)
        legacy = selection_files(index)
        result = academic(index)
        self.assertEqual(
            [r["academic_status"] for r in result["rows"]], ["include", "exclude"]
        )
        self.assertEqual(result["counts"]["exclude"], 1)
        self.assertEqual(result["historical_candidate_total"], "Unknown")
        self.assertEqual(result["input_canonical_sha256"], sha(before))
        self.assertFalse(result["scientific_assessment"])
        self.assertEqual(technical, derive_literature_selection(index))
        self.assertTrue(all(r["status"] == "pending" for r in technical["rows"]))
        self.assertEqual(canonical(index), before)
        exported = selection_files(index, include_academic_screening=True)
        self.assertEqual(len(legacy), 6)
        self.assertEqual({name: exported[name] for name in legacy}, legacy)
        from research_workspace.academic_screening import academic_screening_files

        self.assertEqual(
            {name: raw for name, raw in exported.items() if name not in legacy},
            academic_screening_files(index),
        )
        summary = list(
            csv.DictReader(
                io.StringIO(
                    exported["literature/academic_screening.csv"].decode("utf-8-sig")
                )
            )
        )
        self.assertEqual(
            [(r["work_id"], r["academic_status"]) for r in summary],
            [(r["work_id"], r["academic_status"]) for r in result["rows"]],
        )
        book = load_workbook(
            io.BytesIO(exported["literature/academic_screening.xlsx"]), read_only=True
        )
        values = list(book["AcademicScreening"].iter_rows(values_only=True))
        for column in ("work_id", "current_display_version_id", "academic_status"):
            self.assertEqual(
                [r[values[0].index(column)] for r in values[1:]],
                [r[column] for r in result["rows"]],
            )
        markdown = exported["literature/academic_screening.md"].decode()
        self.assertIn('"exclude": 1', markdown)
        self.assertTrue(all(r["work_id"] in markdown for r in result["rows"]))

    def test_empty_index_keeps_historical_total_and_history_completeness_unknown(self):
        index = fixture_index()
        index.update(papers=[], claims=[], sources=[])
        index["bibliography"].update(entries=[], all_bibtex=bibtex(index).decode())
        result = academic(index)
        self.assertEqual(result["rows"], [])
        self.assertEqual(result["historical_candidate_total"], "Unknown")
        self.assertEqual(result["decision_history_completeness"], "Unknown")

    def test_append_order_reversals_and_unknown_history(self):
        index = fixture_index()
        index["screening"] = [decision(9), decision(10, status="exclude")]
        result = academic(index)
        self.assertEqual(result["rows"][0]["academic_status"], "exclude")
        self.assertEqual(result["rows"][1]["academic_status"], "Unknown")
        self.assertEqual(result["rows"][1]["decision_event_count"], "Unknown")
        self.assertEqual(
            [e["original_event"]["decision_id"] for e in result["events"]],
            ["e9", "e10"],
        )
        self.assertEqual(result["counts"]["unknown"], 1)
        self.assertEqual(result["decision_history_completeness"], "Unknown")

    def test_markdown_like_legacy_identity_fields_remain_literal(self):
        index = fixture_index()
        event = decision(
            1, status="![image](https://example.invalid/image.png)\n# forged"
        )
        event["decision_id"] = "[link](https://example.invalid/)"
        index["screening"] = [event]
        exported = selection_files(index, include_academic_screening=True)
        rendered = MarkdownIt("commonmark", {"html": True}).render(
            exported["literature/academic_screening.md"].decode()
        )
        self.assertNotIn("<a href=", rendered)
        self.assertNotIn("<img ", rendered)
        self.assertNotIn("<h1>forged", rendered)
        self.assertIn(event["decision_id"], rendered)

    def test_exports_escape_text_preserve_templates_and_split_long_history(self):
        index = fixture_index()
        reason = "=template <script>alert(1)</script> | reason\n原样"
        index["screening"] = [
            decision(n, status="exclude", reason=reason) for n in range(500)
        ]
        exported = selection_files(index, include_academic_screening=True)
        result = json.loads(exported["literature/academic_screening.json"])
        self.assertEqual(result["rows"][0]["reason_recorded"], True)
        self.assertEqual(result["rows"][0]["reason_nonempty"], True)
        self.assertEqual(result["events"][0]["original_event"]["reason"], reason)
        history = list(
            csv.DictReader(
                io.StringIO(
                    exported["literature/academic_screening_history.csv"].decode(
                        "utf-8-sig"
                    )
                )
            )
        )
        self.assertEqual(len(history), 500)
        self.assertEqual(history[0]["reason"], "'" + reason)
        md = exported["literature/academic_screening.md"].decode()
        self.assertNotIn("<script>", md)
        self.assertIn("&lt;script&gt;", md)
        book = load_workbook(
            io.BytesIO(exported["literature/academic_screening.xlsx"]), read_only=True
        )
        self.assertEqual(book["ScreeningHistory"].max_row, 501)
        self.assertEqual(book["AcademicScreening"].max_row, 3)
        values = list(book["ScreeningHistory"].iter_rows(values_only=True))
        self.assertEqual(values[1][values[0].index("reason")], reason)
        self.assertEqual(
            exported, selection_files(index, include_academic_screening=True)
        )

    def test_large_individual_reason_uses_complete_json_reference_in_excel(self):
        index = fixture_index()
        index["screening"] = [decision(1, reason="长" * 40000)]
        exported = selection_files(index, include_academic_screening=True)
        result = json.loads(exported["literature/academic_screening.json"])
        self.assertEqual(len(result["events"][0]["original_event"]["reason"]), 40000)
        book = load_workbook(
            io.BytesIO(exported["literature/academic_screening.xlsx"]), read_only=True
        )
        values = list(book["ScreeningHistory"].iter_rows(values_only=True))
        ref = json.loads(values[1][values[0].index("reason")])
        self.assertEqual(ref["path"], "literature/academic_screening.json")
        self.assertEqual(
            ref["sha256"], sha(exported["literature/academic_screening.json"])
        )

    def test_canonical12_unread_library_retains_original_work_level_events(self):
        from test_stage1_screening_records import ScreeningRecordsTests as fixture
        from research_workspace.projection import _project_records

        fixture.setUpClass()
        self.addCleanup(fixture.doClassCleanups)
        records = deepcopy(fixture.records)
        index = _project_records(
            "synthetic-screening",
            "0" * 64,
            {},
            None,
            {"original_input_sha256": sha(canonical(records))},
            records,
            deepcopy(records["sources"]),
            "0" * 64,
            {"execution_authorized": False},
        )
        before = canonical(index)
        result = academic(index)
        library = records["screening_library"]["rows"]
        self.assertEqual(len(result["rows"]), len(library))
        self.assertEqual(
            [e["original_event"] for e in result["events"]],
            [e for row in library for e in row["decision_history"]],
        )
        self.assertTrue(all(e["decision_scope"] == "work" for e in result["events"]))
        self.assertTrue(all(e["decision_version_id"] is None for e in result["events"]))
        self.assertEqual(result["counts"]["not_screened"], 1)
        self.assertEqual(result["count_scope"], "bound-run")
        self.assertEqual(
            result["library_counts"], records["screening_library"]["counts"]
        )
        self.assertEqual(result["library_counts"]["total_candidates"], 5)
        self.assertTrue(any(r["saved_counts"]["reversals"] for r in result["rows"]))
        self.assertEqual(
            result["decision_history_completeness"], "complete-for-bound-library"
        )
        self.assertEqual(canonical(index), before)


if __name__ == "__main__":
    unittest.main()
