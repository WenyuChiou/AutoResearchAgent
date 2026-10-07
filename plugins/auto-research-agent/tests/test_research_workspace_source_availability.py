"""Neutral regressions for source availability and claim separation."""

from copy import deepcopy
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "cli"))

from research_workspace.source_availability import derive_source_availability  # noqa: E402
from research_workspace.source_rerun import attach_rerun, rerun_files  # noqa: E402
from research_workspace.wiki import wiki_files  # noqa: E402
from stage1_deliverable.common import canonical, sha  # noqa: E402
from test_research_workspace_source_rerun_exports import (  # noqa: E402
    fixture_index,
    fixture_manifest,
    write_manifest,
)


class SourceAvailabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="source-availability-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.index = fixture_index()

    def attach(self, manifest=None, root=None):
        manifest = manifest or fixture_manifest(self.index)
        root = root or self.root
        expected = write_manifest(root, manifest)
        return attach_rerun(self.index, root, expected)

    def test_full_text_does_not_prove_claim_and_exports_reconcile(self):
        projected = self.attach()
        value = derive_source_availability(projected)
        self.assertEqual(
            value["counts"],
            {
                "full-text-readable": 1,
                "abstract-only": 1,
                "bibliography-only": 0,
                "engineering-read-failure": 0,
                "historical_bibliography": 2,
                "readable": 2,
            },
        )
        self.assertEqual(value["claim_rows"][0]["original_assessment"], "Unknown")
        self.assertTrue(value["claim_rows"][0]["source_readable"])
        files = rerun_files(projected, self.root)
        self.assertEqual(
            json.loads(files["source-rerun/source-availability.json"]), value
        )
        workbook = load_workbook(
            io.BytesIO(files["source-rerun/catalog.xlsx"]), read_only=True
        )
        self.assertTrue(
            {"SourceAvailability", "ClaimAvailability"} <= set(workbook.sheetnames)
        )
        self.assertEqual(
            len(
                files["source-rerun/source-availability.csv"]
                .decode("utf-8-sig")
                .splitlines()
            ),
            3,
        )
        self.assertIn(
            b"availability_is_not_claim_assessment",
            files["source-rerun/source-availability.md"],
        )
        note = next(
            raw
            for name, raw in wiki_files(projected).items()
            if name != "wiki/README.md" and b'"work_id": "fixture-0"' in raw
        )
        self.assertIn(b"Readable-source status does not change", note)

    def test_metadata_text_is_not_completed_reading_and_failure_is_distinct(self):
        for status, expected in (
            ("metadata-text", "bibliography-only"),
            ("failed-engineering", "engineering-read-failure"),
        ):
            with self.subTest(status=status):
                manifest = fixture_manifest(self.index)
                row = manifest["data"]["rows"][0]
                if status == "metadata-text":
                    row["reading"]["evidence_level"] = "metadata"
                else:
                    manifest["files"].pop(row.pop("extracted_path"))
                    row["reading"].update(
                        status="failed-engineering",
                        evidence_level="metadata",
                        characters=0,
                        text_sha256=None,
                        locators=[],
                        error={"type": "ValueError", "message": "synthetic failure"},
                    )
                with tempfile.TemporaryDirectory(
                    prefix="availability-state-"
                ) as folder:
                    projected = self.attach(manifest, Path(folder).resolve())
                first = derive_source_availability(projected)["source_rows"][0]
                self.assertEqual(first["availability"], expected)
                self.assertFalse(first["readable"])

    def test_accepted_passage_requires_exact_work_source_version_and_raw_hash(self):
        manifest = fixture_manifest(self.index)
        manifest["data"]["rows"][0]["reading"]["evidence_level"] = "metadata"
        projected = self.attach(manifest)
        row = projected["source_rerun"]["data"]["rows"][0]

        def add_passage(**changes):
            passage = {
                "passage_id": "neutral-passage",
                "work_id": row["work_id"],
                "version_id": row["version_id"],
                "source_id": row["source_id"],
                "raw_sha256": row["raw_sha256"],
                "quote": "A neutral saved abstract sentence.",
                "raw_html": {"selector": "main p", "sha256": row["raw_sha256"]},
                "clean_source": {
                    "sha256": sha(b"A neutral saved abstract sentence."),
                    "start": 0,
                    "end": len("A neutral saved abstract sentence."),
                },
                **changes,
            }
            payload = {
                "source_kind": "publication abstract, not full text",
                "passages": [passage],
            }
            text = json.dumps(payload, sort_keys=True)
            changed = deepcopy(projected)
            changed["supplement"] = {
                "status": "accepted-scoped-private",
                "data": {
                    "atomic_revisions": [
                        {
                            "atom_id": "neutral-atomic-revision",
                            "parent_claim_id": "historical-claim",
                            "work_id": row["work_id"],
                            "assessment": "unproved",
                        }
                    ]
                },
                "evidence_files": {
                    "neutral/passages.json": {
                        "sha256": sha(text.encode()),
                        "text": text,
                    },
                    "neutral/abstract.txt": {
                        "sha256": sha(b"A neutral saved abstract sentence."),
                        "text": "A neutral saved abstract sentence.",
                    },
                },
            }
            return derive_source_availability(changed)

        accepted = add_passage()
        self.assertEqual(accepted["source_rows"][0]["availability"], "abstract-only")
        self.assertEqual(
            accepted["claim_rows"][-1]["source_availability"], "version-unmapped"
        )
        self.assertIsNone(accepted["claim_rows"][-1]["version_id"])
        for changes in (
            {"work_id": "other-work"},
            {"version_id": "other-version"},
            {"source_id": "other-source"},
            {"raw_sha256": "0" * 64},
            {"quote": ""},
            {"raw_html": {}},
            {"clean_source": {}},
            {
                "clean_source": {
                    "sha256": "0" * 64,
                    "start": 0,
                    "end": len("A neutral saved abstract sentence."),
                }
            },
            {"raw_html": {"selector": "main p", "sha256": "0" * 64}},
        ):
            with self.subTest(changes=changes):
                self.assertEqual(
                    add_passage(**changes)["source_rows"][0]["availability"],
                    "bibliography-only",
                )

        bound = deepcopy(projected)
        text = json.dumps(
            {
                "atomic_claims": [
                    {
                        "claim_id": "neutral-atom",
                        "work_id": row["work_id"],
                        "version_id": row["version_id"],
                        "source_id": row["source_id"],
                    }
                ],
                "source_binding": {"source_sha256": row["raw_sha256"]},
            }
        )
        bound["supplement"] = {
            "status": "accepted-scoped-private",
            "data": {
                "atomic_revisions": [
                    {
                        "atom_id": "neutral-atom",
                        "parent_claim_id": "historical",
                        "work_id": row["work_id"],
                        "assessment": "unproved",
                        "evidence_ref": "../neutral/atoms.json",
                    }
                ]
            },
            "evidence_files": {
                "neutral/atoms.json": {"text": text, "sha256": sha(text.encode())}
            },
        }
        atom = derive_source_availability(bound)["claim_rows"][-1]
        self.assertEqual(atom["version_id"], row["version_id"])
        self.assertEqual(atom["original_assessment"], "unproved")
        payload = json.loads(text)
        payload["source_binding"]["source_sha256"] = "0" * 64
        bad = json.dumps(payload)
        bound["supplement"]["evidence_files"]["neutral/atoms.json"] = {
            "text": bad,
            "sha256": sha(bad.encode()),
        }
        self.assertEqual(
            derive_source_availability(bound)["claim_rows"][-1]["source_availability"],
            "version-unmapped",
        )

        # A declared version never substitutes for source-bound evidence.
        bound["supplement"]["data"]["atomic_revisions"][0]["version_id"] = row[
            "version_id"
        ]
        for evidence_files in (
            {},
            {"neutral/atoms.json": {"text": bad, "sha256": sha(bad.encode())}},
        ):
            with self.subTest(explicit_version_invalid_evidence=evidence_files):
                invalid = deepcopy(bound)
                invalid["supplement"]["evidence_files"] = evidence_files
                result = derive_source_availability(invalid)["claim_rows"][-1]
                self.assertIsNone(result["version_id"])
                self.assertEqual(result["source_availability"], "version-unmapped")
                self.assertIsNone(result["source_readable"])
                self.assertEqual(result["original_assessment"], "unproved")
        bound["supplement"]["evidence_files"] = {
            "neutral/atoms.json": {"text": text, "sha256": sha(text.encode())}
        }
        self.assertEqual(
            derive_source_availability(bound)["claim_rows"][-1]["version_id"],
            row["version_id"],
        )

        wrong_kind = deepcopy(projected)
        payload = {
            "source_kind": "unclassified saved page",
            "passages": [
                {
                    "passage_id": "neutral-passage",
                    "work_id": row["work_id"],
                    "version_id": row["version_id"],
                    "source_id": row["source_id"],
                    "raw_sha256": row["raw_sha256"],
                    "quote": "A neutral saved sentence.",
                    "raw_html": {"selector": "main p"},
                    "clean_source": {"sha256": "1" * 64, "start": 0, "end": 25},
                }
            ],
        }
        text = json.dumps(payload, sort_keys=True)
        wrong_kind["supplement"] = {
            "status": "accepted-scoped-private",
            "data": {"atomic_revisions": []},
            "evidence_files": {
                "neutral/passages.json": {"sha256": sha(text.encode()), "text": text}
            },
        }
        self.assertEqual(
            derive_source_availability(wrong_kind)["source_rows"][0]["availability"],
            "bibliography-only",
        )

    def test_derivation_is_deterministic_and_does_not_mutate_index(self):
        projected = self.attach()
        before = canonical(projected)
        self.assertEqual(
            canonical(derive_source_availability(projected)),
            canonical(derive_source_availability(projected)),
        )
        self.assertEqual(canonical(projected), before)


if __name__ == "__main__":
    unittest.main()
