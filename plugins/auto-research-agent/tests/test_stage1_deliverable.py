"""Offline production exporter tests with real public source parser/CLI replay."""

import copy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import bibtexparser
from bibtexparser.latexenc import latex_to_unicode
from openpyxl import load_workbook
from requests import Timeout
from requests.structures import CaseInsensitiveDict
from research_hub.source_fetch import fetch_public_source

PLUGIN = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PLUGIN / "cli"))

from stage1_deliverable import package, sources, views  # noqa: E402
from stage1_deliverable.common import (  # noqa: E402
    DeliverableError,
    canonical,
    inventory,
    preflight,
    read_json,
    sha,
    safe_path,
    write_json,
)


TITLE = "Synthetic household research"
HTML = (
    f"<html><head><title>{TITLE}</title></head><body><article><h1>{TITLE}</h1><h2>Introduction</h2><p>"
    + "Synthetic research asks how a declared household scenario changes consumption. "
    * 6
    + "</p><h2>Results</h2><p>"
    + "Synthetic households alter consumption under a declared scenario. " * 12
    + "</p></article></body></html>"
).encode()


def full_source_pdf():
    """An authored three-page article, not a one-line parser smoke probe."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 5 0 R 7 0 R] /Count 3 >>",
    ]
    for number, heading in enumerate(("Introduction", "Methods", "Results")):
        stream = (
            f"BT /F1 12 Tf 50 750 Td ({TITLE}) Tj 0 -24 Td ({heading}) Tj "
            + "0 -20 Td (Synthetic households change consumption in this authored fixture.) Tj "
            * 12
            + "ET"
        ).encode()
        objects.extend(
            [
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 9 0 R >> >> /Contents {4 + number * 2} 0 R >>".encode(),
                f"<< /Length {len(stream)} >>\nstream\n".encode()
                + stream
                + b"\nendstream",
            ]
        )
    objects.extend(
        [
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            f"<< /Title ({TITLE}) /Author (Synthetic Author) >>".encode(),
        ]
    )
    raw, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(raw)
    raw += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    raw += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    return (
        raw
        + f"trailer\n<< /Size {len(offsets)} /Root 1 0 R /Info 10 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )


class Response:
    def __init__(self, raw, content_type, status):
        self.raw = raw
        self.headers = CaseInsensitiveDict({"Content-Type": content_type})
        self.status_code = status

    def iter_content(self, chunk_size):
        yield self.raw

    def close(self):
        pass


class ResearchDeliverableTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # This suite exercises the real no-Git guard. Require a checked
        # external root before creating fixtures, including the fake home.
        root = Path(tempfile.gettempdir()).resolve()
        try:
            preflight(root)
        except (DeliverableError, OSError) as error:
            raise RuntimeError(
                "Exporter tests require a writable temporary root outside Git. "
                "Set TEMP, TMP and TMPDIR to that root before starting Python; "
                "see stage1_deliverable/README.md, Tests and evidence limits."
            ) from error
        cls.private_test_root = root

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="stage1-deliverable-test-", dir=self.private_test_root
        )
        self.addCleanup(self.temp.cleanup)
        # macOS exposes its temp root through /var -> /private/var. Fixtures
        # use the physical directory; production still rejects linked inputs.
        self.root = Path(self.temp.name).resolve()
        self.inputs = self.root / "input"
        self.inputs.mkdir()
        self.output = self.root / "package"

    def test_empty_diagnostics_preserve_the_legacy_source_receipt(self):
        self.make_records()
        result = read_json(self.inputs / "source/source-fetch-result.json")
        result.pop("diagnostics", None)
        legacy = sources.receipt_digest(result)
        self.assertEqual(legacy, result["receipt_sha256"])
        result["diagnostics"] = {}
        self.assertEqual(sources.receipt_digest(result), legacy)

    def test_nonempty_legacy_diagnostics_replay_against_pinned_sdk(self):
        records = self.make_records()
        path = self.inputs / "source/source-fetch-result.json"
        result = read_json(path)
        result["diagnostics"] = {"body_completeness": "pending", "pages_total": 3}
        self.assertEqual(result["receipt_sha256"], sources.receipt_digest(result))
        write_json(path, result)
        records["sources"][0]["result_sha256"] = sha(path.read_bytes())
        archive = self.root / "diagnostics-archive"
        sources.stage_source(records["sources"][0], self.inputs, archive)
        self.assertEqual((archive / "original.json").read_bytes(), path.read_bytes())
        replayed, _, observation = sources.validate_archive(archive)
        self.assertEqual(replayed, result)
        self.assertEqual(observation["returncode"], 0)
        self.assertFalse(observation["network_acquisition"])

    def test_diagnostics_file_tamper_rejects_before_sdk_replay(self):
        records = self.make_records()
        path = self.inputs / "source/source-fetch-result.json"
        result = read_json(path)
        result["diagnostics"] = {"body_completeness": "pending"}
        write_json(path, result)
        records["sources"][0]["result_sha256"] = sha(path.read_bytes())
        result["diagnostics"]["body_completeness"] = "confirmed"
        write_json(path, result)
        with patch("stage1_deliverable.sources.subprocess.run") as replay:
            with self.assertRaisesRegex(DeliverableError, "source result changed"):
                sources.stage_source(
                    records["sources"][0], self.inputs, self.root / "tampered-archive"
                )
        replay.assert_not_called()

    def test_new_diagnostics_checksum_rejects_under_legacy_sdk_pin(self):
        self.make_records()
        result = read_json(self.inputs / "source/source-fetch-result.json")
        result["diagnostics"] = {"body_completeness": "pending"}
        result["receipt_sha256"] = sha(
            canonical(
                {
                    "schema_version": "source-fetch-result/v1",
                    "request": result["request"],
                    "attempts": [
                        {key: a[key] for key in sources.ATTEMPT_FIELDS}
                        for a in result["attempts"]
                    ],
                    "extracted_text_sha256": result["extracted_text_sha256"],
                    "result": {
                        key: result.get(key)
                        for key in (*sources.RESULT_FIELDS, "diagnostics")
                    },
                }
            )
        )
        self.assertNotEqual(result["receipt_sha256"], sources.receipt_digest(result))
        archive = self.root / "diagnostics-archive"
        archive.mkdir()
        write_json(archive / "original.json", result)
        with patch("stage1_deliverable.sources.subprocess.run") as replay:
            with self.assertRaisesRegex(DeliverableError, "original source receipt"):
                sources.validate_archive(archive)
        replay.assert_not_called()

    def make_records(self, raw=HTML, content_type="text/html", status=200, title=TITLE):
        with (
            patch(
                "research_hub.source_fetch._resolve_host_addresses",
                return_value=("93.184.216.34",),
            ),
            patch(
                "requests.Session.get",
                return_value=Response(raw, content_type, status),
                side_effect=raw if isinstance(raw, Exception) else None,
            ),
        ):
            result = fetch_public_source(
                output_dir=self.inputs / "source",
                url="https://example.org/study",
                title=title,
            ).to_dict()
        source = {
            "source_id": "src1",
            "work_id": "work1",
            "version_id": "v1",
            "result_path": "source/source-fetch-result.json",
            "result_sha256": sha(
                (self.inputs / "source/source-fetch-result.json").read_bytes()
            ),
            "access_note": "Synthetic response authored for offline testing; no real paper downloaded.",
        }
        claims = []
        if result["status"] == "available" and result["extracted_text_path"]:
            text = Path(result["extracted_text_path"]).read_bytes().decode()
            claims = [
                {
                    "claim_id": "claim1",
                    "work_id": "work1",
                    "version_id": "v1",
                    "text": "Synthetic illustrative claim; no scientific endorsement.",
                    "source_id": "src1",
                    "relation": "unverified",
                    "evidence_level": result["evidence_level"],
                    "locator": "characters 0:20",
                    "start": 0,
                    "end": 20,
                    "quote": text[:20],
                }
            ]
        paper = {
            "work_id": "work1",
            "version_id": "v1",
            "title": title,
            "authors": ["Synthetic Author"],
            "year": 2024,
            "venue": "Synthetic venue",
            "doi": None,
            "url": "https://example.org/study",
            "evidence_level": result["evidence_level"]
            if result["status"] == "available"
            else "metadata",
            "source_ids": ["src1"],
            "classification": {
                key: "Explicit synthetic classification"
                for key in (
                    "topic_cluster",
                    "method",
                    "geography",
                    "population",
                    "data_type",
                    "domain",
                )
            },
            "roles": [],
            "findings": {
                key: "Not independently assessed in this synthetic fixture"
                for key in (
                    "question",
                    "data",
                    "method",
                    "main_findings",
                    "limitations",
                    "relevance",
                    "transferability",
                )
            },
            "claim_ids": [c["claim_id"] for c in claims],
        }
        value = {
            "kind": "Stage1ResearchRecords",
            "schema_version": "1.0.0",
            "topic": "Synthetic literature task",
            "as_of": "2026-09-28T00:00:00+00:00",
            "papers": [paper],
            "sources": [source],
            "claims": claims,
            "screening": [
                {
                    "decision_id": "decision1",
                    "work_id": "work1",
                    "version_id": "v1",
                    "status": "pending",
                    "reason": "Synthetic record remains unverified",
                    "query": "synthetic query",
                    "discovery_path": "synthetic fixture, no live search",
                    "observed_at": "2026-09-28T00:00:00+00:00",
                }
            ],
            "coverage": [
                {
                    "need_id": "need1",
                    "description": "Synthetic need",
                    "work_ids": ["work1"],
                    "recent_sweep": "not exercised",
                    "closest_work_check": "not exercised",
                    "unresolved": "scientific relevance unverified",
                    "stop_decision": "continue",
                    "reason": "fixture is not research",
                }
            ],
        }
        self.records = self.inputs / "records.json"
        write_json(self.records, value)
        return value

    def build(self, records=None):
        if records is not None:
            write_json(self.records, records)
        return package.build(self.records, self.output, sha(self.records.read_bytes()))

    def rehash(self):
        path = self.output / "provenance_manifest.json"
        value = read_json(path)
        value["files"] = inventory(self.output)
        write_json(path, value)
        return sha(path.read_bytes())

    def test_public_html_package_and_portable_semantic_replay(self):
        self.make_records()
        report = self.build()
        self.assertEqual(
            report["counts"]["acquired_full_text"], {"pdf": 0, "html": 1, "text": 0}
        )
        self.assertEqual(
            load_workbook(self.output / "literature_catalog.xlsx").sheetnames,
            ["Papers", "Classification", "Findings", "Claims", "Screening", "Coverage"],
        )
        moved = self.root / "moved"
        shutil.copytree(self.output, moved)
        shutil.rmtree(self.inputs)
        self.assertEqual(
            package.validate(moved, report["manifest_sha256"])["status"], "passed"
        )

    def test_public_pdf_is_readable_and_bound(self):
        self.make_records(full_source_pdf(), "application/pdf")
        report = self.build()
        self.assertEqual(report["counts"]["acquired_full_text"]["pdf"], 1)

    def test_bibtex_special_characters_and_uri_round_trip(self):
        records = self.make_records()
        paper = records["papers"][0]
        paper["title"] = "\\ {literal} 50% A&B under_score #tag 日本\r\nnext"
        paper["authors"] = ["Author & Co", "René Example"]
        paper["venue"] = "Price $5 ~ approximate ^ exponent"
        paper["url"] = "https://example.org/a_b?q=50%25&other=x_y#section"
        paper["doi"] = "10.1234/a_b%25"
        # A second entry after the hazardous field exposes swallowed entries.
        second = copy.deepcopy(paper)
        second.update(
            work_id="work2", title="Unmatched { opening", url="https://example.org/{id}"
        )
        third = copy.deepcopy(paper)
        third.update(work_id="work3", title="Unmatched } closing")
        records["papers"] += [second, third]
        raw = views.bibtex(records).decode("utf-8")
        database = bibtexparser.loads(raw)
        self.assertEqual(
            [e["ID"] for e in database.entries], ["work1", "work2", "work3"]
        )
        entry = database.entries[0]
        self.assertEqual(
            entry["title"],
            r"{\textbackslash} {\textbraceleft}literal{\textbraceright} 50\% A\&B under\_score \#tag 日本  next",
        )
        self.assertEqual(entry["author"], r"Author \& Co and René Example")
        self.assertEqual(
            entry["journal"],
            r"Price \$5 {\textasciitilde} approximate {\textasciicircum} exponent",
        )
        self.assertEqual(entry["url"], paper["url"])
        self.assertEqual(entry["doi"], paper["doi"])
        self.assertEqual(database.entries[1]["url"], "https://example.org/%7Bid%7D")
        self.assertEqual(
            database.entries[1]["title"], r"Unmatched {\textbraceleft} opening"
        )
        self.assertEqual(
            database.entries[2]["title"], r"Unmatched {\textbraceright} closing"
        )
        self.assertEqual(
            bibtexparser.loads(bibtexparser.dumps(database)).entries, database.entries
        )
        # Independent decoding checks text fidelity, beyond accepting .bib syntax.
        paper["title"] = "\\ 50% A&B under_score #tag 日本"
        decoded = latex_to_unicode(
            bibtexparser.loads(views.bibtex({"papers": [paper]}).decode()).entries[0][
                "title"
            ]
        )
        self.assertEqual(decoded, paper["title"])

    def test_home_git_checkout_requires_external_private_and_temp_roots(self):
        home = self.root / "home"
        home.mkdir()
        subprocess.run(["git", "init", str(home)], check=True, capture_output=True)
        inside = home / "Temp"
        outside = self.root / "private"
        safe_temp = self.root / "temp"
        for path in (inside, outside, safe_temp):
            path.mkdir()
        # A real home checkout can also contain the suite's default TEMP.
        # Reject that host configuration before any fixture assumes safety.
        with patch("tempfile.tempdir", str(inside)):
            with self.assertRaisesRegex(RuntimeError, "Set TEMP, TMP and TMPDIR"):
                type(self).setUpClass()
        env = dict(os.environ, PYTHONPATH=str(PLUGIN / "cli"))
        for output, temporary, expected in (
            (inside, safe_temp, 1),
            (outside, inside, 1),
            (outside, safe_temp, 0),
        ):
            with self.subTest(output=output, temporary=temporary):
                env.update(
                    TEMP=str(temporary), TMP=str(temporary), TMPDIR=str(temporary)
                )
                result = subprocess.run(
                    [
                        sys.executable,
                        "-m",
                        "stage1_deliverable",
                        "preflight",
                        str(output),
                    ],
                    env=env,
                    capture_output=True,
                    text=True,
                )
                self.assertEqual(result.returncode, expected, result.stderr)
                if expected:
                    self.assertIn("Git", result.stderr)
                else:
                    self.assertEqual(
                        json.loads(result.stdout),
                        {
                            "status": "passed",
                            "private_root": str(outside),
                            "temporary_root": str(safe_temp),
                            "repository_guard": "outside-git",
                            "write_probe": "passed",
                        },
                    )
        self.assertEqual(list(outside.iterdir()), [])
        self.assertEqual(list(safe_temp.iterdir()), [])
        self.make_records(full_source_pdf(), "application/pdf")
        self.output = outside / "package"
        with patch("tempfile.tempdir", str(safe_temp)):
            report = self.build()
            self.assertEqual(report["status"], "passed")
            self.assertEqual(
                package.validate(self.output, report["manifest_sha256"])["status"],
                "passed",
            )
        # Skipping preflight cannot write a private replay copy under Git.
        with patch("tempfile.tempdir", str(inside)):
            with self.assertRaisesRegex(DeliverableError, "Git"):
                package.validate(self.output, report["manifest_sha256"])
        self.assertEqual(list(inside.iterdir()), [])

    def test_json_key_order_preserves_views_and_original_input_binding(self):
        records = self.make_records()
        self.build()

        def reversed_keys(value):
            if isinstance(value, dict):
                return {key: reversed_keys(value[key]) for key in reversed(value)}
            if isinstance(value, list):
                return [reversed_keys(item) for item in value]
            return value

        original = json.dumps(reversed_keys(records), indent=2).encode("utf-8")
        self.records.write_bytes(original)
        reordered = self.root / "reordered"
        report = package.build(self.records, reordered, sha(original))
        self.assertEqual(report["status"], "passed")
        self.assertEqual((reordered / "records.original.json").read_bytes(), original)
        self.assertEqual(report["records_sha256"], sha(original))
        for name in (
            "literature_catalog.xlsx",
            "literature_review.md",
            "literature_review.docx",
            "references.bib",
            "claims_and_evidence.csv",
            "search_and_screening.csv",
            "coverage_and_stop.md",
            "README.md",
            "papers.jsonl",
            "paper_manifest.jsonl",
        ):
            with self.subTest(name=name):
                self.assertEqual(
                    (self.output / name).read_bytes(), (reordered / name).read_bytes()
                )

    def test_paywall_unavailable_and_login_disguised_as_pdf_remain_distinct(self):
        cases = [
            (b"<html>purchase this article</html>", "text/html", 402, "paywalled"),
            (b"<html>sign in to continue</html>", "application/pdf", 200, "login-page"),
            (b"not found", "text/plain", 404, "not-found"),
            (b"bad PDF", "application/pdf", 200, "parse-error"),
            (b"slow down", "text/plain", 429, "rate-limited"),
            (b"service unavailable", "text/plain", 503, "network-error"),
            (Timeout("synthetic offline timeout"), "", None, "network-error"),
            (
                f"<html><title>{TITLE}</title><body>Citation metadata only.</body></html>".encode(),
                "text/html",
                200,
                "metadata-only",
            ),
            (
                HTML.replace(TITLE.encode(), b"An entirely unrelated source identity"),
                "text/html",
                200,
                "identity-mismatch",
            ),
        ]
        for raw, mime, code, expected in cases:
            with self.subTest(expected=expected):
                if (self.inputs / "source").exists():
                    shutil.rmtree(self.inputs / "source")
                if self.output.exists():
                    shutil.rmtree(self.output)
                self.make_records(raw, mime, code)
                report = self.build()
                self.assertEqual(report["counts"]["access_state_counts"][expected], 1)
                self.assertEqual(list((self.output / "papers").iterdir()), [])
                if code in (None, 503, 429, 404, 402):
                    import json

                    attempt = json.loads(
                        (self.output / "paper_manifest.jsonl")
                        .read_text(encoding="utf-8")
                        .splitlines()[0]
                    )
                    self.assertFalse(attempt["parser_attempted"])
                    self.assertIsNone(attempt["parser"])

    def test_wrong_work_and_version_rejected(self):
        records = self.make_records()
        for key, value in (("work_id", "other"), ("version_id", "v2")):
            bad = copy.deepcopy(records)
            bad["sources"][0][key] = value
            with (
                self.subTest(key=key),
                self.assertRaisesRegex(DeliverableError, "work|version"),
            ):
                self.build(bad)

    def test_abstract_as_full_text_rejected(self):
        raw = f'<html><head><title>{TITLE}</title></head><body><div class="abstract">A synthetic abstract describes household consumption without full article content.</div></body></html>'.encode()
        records = self.make_records(raw)
        self.assertEqual(records["papers"][0]["evidence_level"], "abstract")
        records["papers"][0]["evidence_level"] = "full-text"
        with self.assertRaisesRegex(DeliverableError, "cannot become full text"):
            self.build(records)

    def test_tampered_source_and_rehashed_extraction_rejected(self):
        self.make_records()
        self.build()
        archive = self.output / "sources/src1"
        result_path = archive / "original.json"
        result = read_json(result_path)
        relative = sources.artifact_map(result)[result["extracted_text_path"]]
        (archive / relative).write_text("invented evidence", encoding="utf-8")
        result["extracted_text_sha256"] = sha((archive / relative).read_bytes())
        result["receipt_sha256"] = sources.receipt_digest(result)
        write_json(result_path, result)
        manifest = read_json(self.output / "provenance_manifest.json")
        manifest["canonical_records"]["sources"][0]["result_sha256"] = sha(
            result_path.read_bytes()
        )
        write_json(self.output / "records.original.json", manifest["canonical_records"])
        manifest["original_input_sha256"] = sha(
            (self.output / "records.original.json").read_bytes()
        )
        write_json(self.output / "provenance_manifest.json", manifest)
        with self.assertRaisesRegex(DeliverableError, "semantic replay"):
            package.validate(self.output, self.rehash())

    def test_cross_format_id_and_count_drift_rejected_even_rehashed(self):
        self.make_records()
        self.build()
        workbook = load_workbook(self.output / "literature_catalog.xlsx")
        workbook["Papers"]["A2"] = "different-work"
        workbook.save(self.output / "literature_catalog.xlsx")
        with self.assertRaisesRegex(DeliverableError, "canonical-to-view"):
            package.validate(self.output, self.rehash())

    def test_duplicate_ids_and_unsafe_paths_rejected(self):
        records = self.make_records()
        bad = copy.deepcopy(records)
        bad["papers"].append(copy.deepcopy(bad["papers"][0]))
        with self.assertRaisesRegex(DeliverableError, "duplicate"):
            self.build(bad)
        records["sources"][0]["result_path"] = "../outside.json"
        with self.assertRaisesRegex(DeliverableError, "unsafe"):
            self.build(records)

    def test_resume_no_reexecution_and_external_hash_binding(self):
        self.make_records()
        report = self.build()
        before = inventory(self.output)
        with patch.object(sources, "stage_source") as acquire:
            with self.assertRaises(FileExistsError):
                self.build()
            acquire.assert_not_called()
        self.assertEqual(before, inventory(self.output))
        with self.assertRaisesRegex(DeliverableError, "trusted external"):
            package.validate(self.output, "0" * 64)
        self.assertEqual(report["status"], "passed")

    def test_repository_output_and_credentials_rejected(self):
        records = self.make_records()
        with self.assertRaisesRegex(DeliverableError, "Git"):
            package.build(
                self.records,
                PLUGIN / "private-deliverable-test",
                sha(self.records.read_bytes()),
            )
        records["papers"][0]["url"] = "https://user:password@example.org/study"
        with self.assertRaisesRegex(DeliverableError, "credentials"):
            self.build(records)

    def test_claim_text_tamper_rejected(self):
        records = self.make_records()
        records["claims"][0]["quote"] = "invented quotation"
        with self.assertRaisesRegex(DeliverableError, "quote differs"):
            self.build(records)

    def test_runtime_bytes_bound(self):
        self.make_records()
        report = self.build()
        runtime = sources.runtime_binding()
        runtime["installed_distributions"]["pdfplumber"]["files_sha256"] = "0" * 64
        with patch.object(sources, "runtime_binding", return_value=runtime):
            with self.assertRaisesRegex(DeliverableError, "runtime bytes"):
                package.validate(self.output, report["manifest_sha256"])

    def test_dependency_sha_bound(self):
        with patch.object(sources.importlib.metadata, "distribution") as distribution:
            distribution.return_value.read_text.return_value = (
                '{"vcs_info":{"commit_id":"wrong"}}'
            )
            with self.assertRaisesRegex(DeliverableError, "pinned merge commit"):
                sources.runtime_binding()

    def test_out_of_bounds_span_and_invented_locator_rejected(self):
        records = self.make_records()
        bad = copy.deepcopy(records)
        bad["claims"][0]["locator"] = "page 9999"
        with self.assertRaisesRegex(DeliverableError, "locator"):
            self.build(bad)
        result = read_json(self.inputs / "source/source-fetch-result.json")
        text = Path(result["extracted_text_path"]).read_text(encoding="utf-8")
        claim = records["claims"][0]
        claim.update(end=len(text) + 100000, quote=text)
        claim["locator"] = f"characters 0:{claim['end']}"
        with self.assertRaisesRegex(DeliverableError, "quote differs"):
            self.build(records)

    def test_rehashed_historical_validation_report_rejected(self):
        self.make_records()
        self.build()
        path = self.output / "sources/src1/validation.json"
        observation = read_json(path)
        import json

        report = json.loads(observation["stdout"])
        report["valid"] = False
        observation["stdout"] = json.dumps(report)
        write_json(path, observation)
        with self.assertRaisesRegex(DeliverableError, "validation report differs"):
            package.validate(self.output, self.rehash())

    def test_manifest_contract_fields_rejected_even_rehashed(self):
        self.make_records()
        self.build()
        path = self.output / "provenance_manifest.json"
        original = read_json(path)
        for key, value in (
            ("experimental", False),
            ("improvement", "improved"),
            ("docx", {"status": "unavailable"}),
        ):
            changed = copy.deepcopy(original)
            changed[key] = value
            write_json(path, changed)
            with (
                self.subTest(key=key),
                self.assertRaisesRegex(DeliverableError, "contract"),
            ):
                package.validate(self.output, self.rehash())

    def test_json_credential_in_source_rejected_before_archive(self):
        raw = HTML.replace(
            b"</body>",
            b'<script>{"password":"SYNTHETIC_DO_NOT_USE","access_token":"SYNTHETIC_DO_NOT_USE"}</script></body>',
        )
        self.make_records(raw)
        with self.assertRaisesRegex(DeliverableError, "credential-like"):
            self.build()
        self.assertFalse((self.output / "sources/src1/original.json").exists())
        with self.assertRaises(DeliverableError):
            sources.reject_secrets(b'{"\\u0070assword":"SYNTHETIC_DO_NOT_USE"}')

    def test_excel_oversize_cells_rejected_without_truncation(self):
        records = self.make_records()
        records["papers"][0]["findings"]["limitations"] = "x" * 32768
        with self.assertRaisesRegex(DeliverableError, "Excel cell limit"):
            self.build(records)

    def test_spreadsheet_and_csv_formula_content_is_inert(self):
        records = self.make_records()
        records["screening"][0]["reason"] = '=HYPERLINK("https://example.org")'
        self.build(records)
        workbook = load_workbook(
            self.output / "literature_catalog.xlsx", data_only=False
        )
        self.assertTrue(
            all(c.data_type != "f" for sheet in workbook for row in sheet for c in row)
        )
        self.assertIn(
            b"'=HYPERLINK", (self.output / "search_and_screening.csv").read_bytes()
        )

    def test_unknown_receipt_credentials_rejected_before_copy(self):
        records = self.make_records()
        path = self.inputs / "source/source-fetch-result.json"
        receipt = read_json(path)
        receipt["Authorization"] = "Bearer SYNTHETIC_CREDENTIAL_DO_NOT_USE"
        write_json(path, receipt)
        records["sources"][0]["result_sha256"] = sha(path.read_bytes())
        with self.assertRaisesRegex(DeliverableError, "unexpected"):
            self.build(records)
        self.assertFalse((self.output / "sources/src1/original.json").exists())

    def test_windows_reserved_and_alias_paths_rejected(self):
        for path in (
            "papers/CON.txt",
            "papers/foo./x.txt",
            "papers/name /x.txt",
            "papers//x.txt",
            "papers/LPT1.pdf",
        ):
            with self.subTest(path=path), self.assertRaises(DeliverableError):
                safe_path(self.root, path)

    def test_original_input_binding_is_replayed(self):
        self.make_records()
        self.build()
        manifest = read_json(self.output / "provenance_manifest.json")
        manifest["original_input_sha256"] = "0" * 64
        write_json(self.output / "provenance_manifest.json", manifest)
        with self.assertRaisesRegex(DeliverableError, "original input"):
            package.validate(self.output, self.rehash())

    def test_institution_and_login_uri_states(self):
        receipt = {"raw_path": None, "status": "inaccessible"}
        for body, url in (
            (b"access through your institution", "https://example.org/paper"),
            (b"Access required", "https://example.org/login"),
        ):
            attempt = {"http_status": 200, "outcome": "inaccessible", "final_url": url}
            self.assertEqual(sources._state(receipt, attempt, body), "login-page")


if __name__ == "__main__":
    unittest.main()
