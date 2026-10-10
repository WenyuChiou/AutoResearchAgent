"""Synthetic acceptance-root, source/page binding and technical-gate checks."""

from copy import deepcopy
from pathlib import Path
import contextlib
import io
import json
import sys
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parents[1] / "cli"))

from test_research_workspace_body_review import fixture
from research_workspace.body_completeness import assess_body_completeness
from research_workspace.literature_selection import (
    _binding,
    _eligible,
    _selection_workbook_tables,
    _table_rows,
)
from research_workspace.whole_source_review import (
    SCOPE,
    attach_whole_source_reviews,
    review_key,
    source_binding,
    validate_whole_source_reviews,
)
from stage1_deliverable.common import DeliverableError, canonical, sha
from stage1_deliverable.views import workbook_bytes
from openpyxl import load_workbook


def bundle(
    change_row=None,
    change_manifest=None,
    change_acceptance=None,
    representation="extracted-text",
):
    row, artifacts, _, _ = fixture(geometry=True)
    # Explicitly authored fixture evidence; not an actual PDF audit.
    row["reading"]["identity_status"] = "consistent"
    row["reading"]["diagnostics"]["reading_order"] = {"status": "pending"}
    row["original_attempts"] = [{"sequence": 1, "response_truncated": False}]
    row["selected_sequence"] = 1
    if change_row:
        change_row(row)
    binding = source_binding(row)
    fence = {
        "formal_admission": False,
        "claim_assessments_changed": False,
        "official_stage2_import_eligible": False,
        "research_execution": False,
        "quality_score": None,
        "reviewer_identity_authenticated": False,
    }
    text = artifacts[row["extracted_path"]].decode()
    pages = []
    for locator in row["reading"]["locators"]:
        n = locator["value"]
        name = f"review/page-{n}.png"
        artifacts[name] = f"authored synthetic render{n}; not an actual image".encode()
        pages.append(
            {
                "page": n,
                "start": locator["start"],
                "end": locator["end"],
                "text_sha256": sha(text[locator["start"] : locator["end"]].encode()),
                "render_path": name,
                "render_sha256": sha(artifacts[name]),
                "body_checked": True,
                "reading_order_checked": True,
                "fidelity_checked": True,
                "evidence_refs": [name],
                "body_status": "confirmed",
                "reading_order_status": "confirmed",
                "fidelity_status": "confirmed",
                "extracted_fidelity_status": "confirmed",
                "extracted_reading_order_status": "confirmed",
                "limitations": [],
            }
        )
    artifacts["review/render-program.py"] = b"synthetic-renderer"
    artifacts["review/render-runtime.json"] = b"synthetic-runtime"
    proof = {
        "kind": "PDFRenderReproductionCheck",
        "schema_version": "1.0.0",
        "raw_sha256": row["raw_sha256"],
        "pages": [
            {
                "page": p["page"],
                "path": p["render_path"],
                "sha256": p["render_sha256"],
                "matches_raw_render": True,
            }
            for p in pages
        ],
        "renderer": {
            "name": "synthetic",
            "version": "synthetic-v1",
            "program_path": "review/render-program.py",
            "runtime_path": "review/render-runtime.json",
            "program_sha256": sha(b"synthetic-renderer"),
            "runtime_sha256": sha(b"synthetic-runtime"),
        },
    }
    artifacts["review/render-proof.json"] = canonical(proof)
    manifest = {
        "kind": "Stage1WholeSourceReview",
        "schema_version": "1.0.0",
        "scope": SCOPE,
        "representation": representation,
        "render_proof": {
            "path": "review/render-proof.json",
            "sha256": sha(canonical(proof)),
        },
        "binding": binding,
        **fence,
        "reviewer_id": "synthetic-independent-reviewer",
        "page_reviews": pages,
        "evidence_refs": [row["raw_path"], row["extracted_path"]],
        "artifact_bindings": {name: sha(raw) for name, raw in artifacts.items()},
        "unresolved_material_defects": [],
        "limitations": [],
    }
    if change_manifest:
        change_manifest(manifest)
    raw = canonical(manifest)
    artifacts["review/manifest.json"] = raw
    acceptance = {
        "kind": "Stage1WholeSourceReviewAcceptance",
        "schema_version": "1.0.0",
        "scope": SCOPE,
        "binding": binding,
        **fence,
        "manifest_sha256": sha(raw),
        "decision": "accepted",
        "reviewer_id": "synthetic-parent-acceptance",
        "unresolved_material_defects": [],
        "evidence_refs": [row["raw_path"]],
    }
    if change_acceptance:
        change_acceptance(acceptance)
    artifacts["review/acceptance.json"] = canonical(acceptance)
    expected = sha(artifacts["review/acceptance.json"])
    row["whole_source_review"] = {
        "manifest_path": "review/manifest.json",
        "manifest_sha256": sha(raw),
        "acceptance_path": "review/acceptance.json",
        "acceptance_sha256": expected,
    }
    hashes = {
        name: {"sha256": sha(raw), "bytes": len(raw)} for name, raw in artifacts.items()
    }
    return row, artifacts, hashes, expected


class WholeSourceReviewTests(unittest.TestCase):
    def test_long_page_audit_exports_complete_checks_without_cell_overflow(self):
        row, artifacts, hashes, digest = bundle()
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        binding = _binding(row, hashes, source_reviews=cache)
        binding["reviewed_body_completeness"]["extracted_page_checks"] = [
            {
                "page": n,
                "fidelity_status": "partial",
                "reading_order_status": "confirmed",
                "limitations": [
                    f"Page {n}: synthetic unresolved sign; " + "detail " * 60
                ],
            }
            for n in range(1, 95)
        ]
        selection = {
            "rows": [
                {
                    "work_id": row["work_id"],
                    "version_id": row["version_id"],
                    "source_ids": [row["source_id"]],
                    "eligible_source_ids": [],
                    "reasons": [],
                    "source_binding_references": [binding],
                }
            ]
        }
        before = canonical(selection)
        with self.assertRaisesRegex(DeliverableError, "Excel cell limit"):
            workbook_bytes({"Selection": _table_rows(selection["rows"])})
        tables = _selection_workbook_tables(selection)
        book = load_workbook(io.BytesIO(workbook_bytes(tables)), read_only=True)
        self.assertEqual(book["SourceReviewPages"].max_row, 95)
        headers, values = list(book["Selection"].values)
        reference = json.loads(dict(zip(headers, values))["source_binding_references"])
        self.assertEqual(reference["sha256"], sha(before))
        self.assertEqual(reference["path"], "literature/selection.json")
        self.assertFalse(reference["content_truncated"])
        self.assertEqual(canonical(selection), before)
        self.assertEqual(
            len(
                json.loads(
                    _table_rows(selection["rows"])[0]["source_binding_references"]
                )[0]["reviewed_body_completeness"]["extracted_page_checks"]
            ),
            94,
        )
        self.assertEqual(
            tables["SourceReviewPages"][-1]["limitations"],
            json.dumps(
                binding["reviewed_body_completeness"]["extracted_page_checks"][-1][
                    "limitations"
                ]
            ),
        )
        book.close()

    def test_single_oversized_page_limit_remains_rejected(self):
        row, artifacts, hashes, digest = bundle()
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        binding = _binding(row, hashes, source_reviews=cache)
        binding["reviewed_body_completeness"]["extracted_page_checks"][0][
            "limitations"
        ] = ["x" * 32768]
        selection = {
            "rows": [
                {
                    "work_id": row["work_id"],
                    "version_id": row["version_id"],
                    "source_ids": [],
                    "eligible_source_ids": [],
                    "reasons": [],
                    "source_binding_references": [binding],
                }
            ]
        }
        with self.assertRaisesRegex(DeliverableError, "Excel cell limit"):
            workbook_bytes(_selection_workbook_tables(selection))

    def test_binding_reference_limit_counts_utf16_units(self):
        row, artifacts, hashes, digest = bundle()
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        binding = _binding(row, hashes, source_reviews=cache)
        binding["error"] = "\U0001f600" * 18000
        selection = {
            "rows": [
                {
                    "work_id": row["work_id"],
                    "version_id": row["version_id"],
                    "source_ids": [],
                    "eligible_source_ids": [],
                    "reasons": [],
                    "source_binding_references": [binding],
                }
            ]
        }
        original = _table_rows(selection["rows"])[0]["source_binding_references"]
        self.assertLess(len(original), 32767)
        tables = _selection_workbook_tables(selection)
        reference = json.loads(tables["Selection"][0]["source_binding_references"])
        self.assertEqual(reference["sha256"], sha(canonical(selection)))
        workbook_bytes(tables)

    def test_cli_requires_rerun_and_forwards_external_whole_source_pin(self):
        from research_workspace.__main__ import main

        argv = [
            "research_workspace",
            "package",
            "output",
            "--project-id",
            "synthetic-project",
            "--expected-manifest-sha256",
            "0" * 64,
            "--reference-root",
            "references",
            "--expected-whole-source-review-acceptance-sha256",
            "1" * 64,
        ]
        output = io.StringIO()
        with (
            patch.object(sys, "argv", argv),
            patch("research_workspace.__main__.project_package", return_value={}),
            contextlib.redirect_stderr(output),
        ):
            self.assertEqual(main(), 1)
        self.assertIn(
            "Whole source review acceptance requires",
            json.loads(output.getvalue())["reason"],
        )
        argv += [
            "--source-rerun-root",
            "rerun",
            "--expected-source-rerun-manifest-sha256",
            "2" * 64,
        ]
        with (
            patch.object(sys, "argv", argv),
            patch("research_workspace.__main__.project_package", return_value={}),
            patch(
                "research_workspace.source_rerun.attach_rerun", return_value={}
            ) as attach,
            patch("research_workspace.__main__.write_workspace", return_value={}),
            contextlib.redirect_stdout(io.StringIO()),
        ):
            self.assertEqual(main(), 0)
        self.assertEqual(
            attach.call_args.kwargs["expected_whole_source_review_acceptance_hashes"],
            ["1" * 64],
        )

    def test_rendered_source_reading_keeps_text_defects_visible(self):
        def partial(m):
            m["page_reviews"][0].update(
                extracted_fidelity_status="failed",
                extracted_reading_order_status="partial",
                limitations=[
                    "Synthetic negative sign unresolved in extracted text; read raw page; no text-only claim support"
                ],
            )
            m["limitations"] = [
                "Extracted text fidelity is partial; complete raw source page reading was reviewed"
            ]

        row, artifacts, hashes, digest = bundle(
            change_manifest=partial, representation="rendered-raw-with-extracted-text"
        )
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        result = _binding(row, hashes, source_reviews=cache)
        self.assertTrue(_eligible(row, cache))
        self.assertEqual(result["body_completeness"]["status"], "pending")
        self.assertEqual(
            result["reviewed_body_completeness"]["representation"],
            "rendered-raw-with-extracted-text",
        )
        self.assertEqual(
            cache[review_key(row)]["manifest"]["page_reviews"][0][
                "extracted_fidelity_status"
            ],
            "failed",
        )
        for representation, limits in (
            ("extracted-text", ["failure"]),
            ("rendered-raw-with-extracted-text", []),
        ):

            def undisclosed(m):
                m["page_reviews"][0].update(
                    extracted_fidelity_status="failed", limitations=limits
                )

            row, artifacts, hashes, digest = bundle(
                change_manifest=undisclosed, representation=representation
            )
            with self.assertRaises(DeliverableError):
                attach_whole_source_reviews([row], artifacts, hashes, [digest])

    def test_page_only_extraction_defects_reach_selection_summary(self):
        limitation = "Synthetic negative sign lost; use raw page"

        def partial(m):
            m["page_reviews"][0].update(
                extracted_fidelity_status="failed",
                extracted_reading_order_status="partial",
                limitations=[limitation],
            )

        row, artifacts, hashes, digest = bundle(
            change_manifest=partial, representation="rendered-raw-with-extracted-text"
        )
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        self.assertEqual(cache[review_key(row)]["manifest"]["limitations"], [])
        summary = _binding(row, hashes, source_reviews=cache)[
            "reviewed_body_completeness"
        ]
        self.assertTrue(_eligible(row, cache))
        self.assertEqual(summary["status"], "confirmed")
        self.assertEqual(summary["limitations"], [f"Page 1: {limitation}"])
        self.assertEqual(
            summary["extracted_page_checks"][0],
            {
                "page": 1,
                "fidelity_status": "failed",
                "reading_order_status": "partial",
                "limitations": [limitation],
            },
        )
        self.assertEqual(
            summary["extracted_page_checks"][1]["fidelity_status"], "confirmed"
        )
        self.assertEqual(summary["extracted_page_checks"][1]["limitations"], [])
        self.assertEqual(
            _binding(row, hashes, source_reviews=cache)["reviewed_body_completeness"],
            summary,
        )

    def test_raw_render_proof_must_cover_current_source_and_all_pages(self):
        row, artifacts, hashes, digest = bundle()
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        for change in ("raw", "pages", "failed", "runtime"):
            candidate = deepcopy(cache)
            proof = candidate[review_key(row)]["render_proof"]
            if change == "raw":
                proof["raw_sha256"] = "0" * 64
            elif change == "pages":
                proof["pages"].pop()
            elif change == "failed":
                proof["pages"][0]["matches_raw_render"] = False
            else:
                proof["renderer"]["version"] = ""
            with self.subTest(change=change), self.assertRaises(DeliverableError):
                validate_whole_source_reviews(
                    {
                        "data": {"rows": [row]},
                        "artifact_hashes": hashes,
                        "whole_source_reviews": candidate,
                    }
                )

    def test_accepted_review_changes_only_effective_body_gate(self):
        row, artifacts, hashes, digest = bundle()
        before = deepcopy(row)
        self.assertFalse(_eligible(row))
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        self.assertEqual(row, before)
        extension = {
            "data": {"rows": [row]},
            "artifact_hashes": hashes,
            "whole_source_reviews": cache,
        }
        self.assertEqual(validate_whole_source_reviews(extension), cache)
        self.assertTrue(_eligible(row, cache))
        self.assertEqual(assess_body_completeness(row)["status"], "pending")
        result = _binding(row, hashes, source_reviews=cache)
        self.assertEqual(result["body_completeness"]["status"], "pending")
        self.assertEqual(result["reviewed_body_completeness"]["status"], "confirmed")
        self.assertFalse(cache[review_key(row)]["manifest"]["formal_admission"])

    def test_no_carrier_preserves_legacy_and_unused_pins_reject(self):
        self.assertEqual(attach_whole_source_reviews([{}], {}, {}, None), {})
        self.assertEqual(
            validate_whole_source_reviews(
                {"data": {"rows": [{}]}, "artifact_hashes": {}}
            ),
            {},
        )
        with self.assertRaisesRegex(DeliverableError, "unused external"):
            attach_whole_source_reviews([], {}, {}, ["0" * 64])

    def test_external_acceptance_is_required_and_duplicates_reject(self):
        row, artifacts, hashes, digest = bundle()
        for pins in ([], ["0" * 64], [digest, digest]):
            with self.subTest(pins=pins), self.assertRaises(DeliverableError):
                attach_whole_source_reviews([row], artifacts, hashes, pins)

    def test_incomplete_source_identity_or_acquisition_cannot_be_overridden(self):
        changes = [
            lambda r: r["reading"].__setitem__("identity_status", "unverified"),
            lambda r: r["original_attempts"][0].__setitem__("response_truncated", True),
            lambda r: r["reading"]["diagnostics"].__setitem__("omitted_pages", [2]),
            lambda r: r["reading"]["diagnostics"].__setitem__("blank_pages", [2]),
            lambda r: r["reading"]["diagnostics"].__setitem__("truncated", True),
            lambda r: r["reading"]["diagnostics"].__setitem__(
                "omitted_words", ["lost"]
            ),
            lambda r: r["reading"]["diagnostics"].__setitem__(
                "extraction_fidelity", "incomplete"
            ),
            lambda r: r["reading"].__setitem__("evidence_level", "abstract"),
            lambda r: r["reading"]["diagnostics"].pop("pages_total"),
        ]
        for i, change in enumerate(changes):
            with self.subTest(case=i):
                row, artifacts, hashes, digest = bundle(change_row=change)
                with self.assertRaises(DeliverableError):
                    attach_whole_source_reviews([row], artifacts, hashes, [digest])

    def test_every_page_check_extent_render_and_reference_are_required(self):
        changes = [
            lambda m: m["page_reviews"].pop(),
            lambda m: m["page_reviews"].reverse(),
            lambda m: m["page_reviews"][1].__setitem__("page", 1),
            lambda m: m["page_reviews"][0].__setitem__("fidelity_checked", False),
            lambda m: m["page_reviews"][0].__setitem__("body_checked", 1),
            lambda m: m["page_reviews"][0].__setitem__("fidelity_status", "failed"),
            lambda m: m["page_reviews"][0].__setitem__(
                "reading_order_status", "pending"
            ),
            lambda m: m["page_reviews"][0].__setitem__("evidence_refs", []),
            lambda m: m["page_reviews"][0].__setitem__("start", 1),
            lambda m: m["page_reviews"][0].__setitem__("start", False),
            lambda m: m["page_reviews"][0].__setitem__("text_sha256", "0" * 64),
            lambda m: m["page_reviews"][0].__setitem__("render_sha256", "0" * 64),
            lambda m: m["evidence_refs"].append("unbound.json"),
            lambda m: m.__setitem__(
                "unresolved_material_defects", ["lost negative sign"]
            ),
            lambda m: m.__setitem__("scope", "engineering-only-body-review"),
        ]
        for i, change in enumerate(changes):
            with self.subTest(case=i):
                row, artifacts, hashes, digest = bundle(change_manifest=change)
                with self.assertRaises(DeliverableError):
                    attach_whole_source_reviews([row], artifacts, hashes, [digest])

    def test_promotion_or_unaccepted_manifest_rejects(self):
        for change in (
            lambda a: a.__setitem__("official_stage2_import_eligible", True),
            lambda a: a.__setitem__("formal_admission", True),
            lambda a: a.__setitem__("quality_score", 90),
            lambda a: a.__setitem__("manifest_sha256", "0" * 64),
            lambda a: a.__setitem__("decision", "pending"),
            lambda a: a.__setitem__("reviewer_id", ""),
        ):
            row, artifacts, hashes, digest = bundle(change_acceptance=change)
            with self.assertRaises(DeliverableError):
                attach_whole_source_reviews([row], artifacts, hashes, [digest])

    def test_changed_artifact_or_cache_rejects(self):
        for label in (
            "raw",
            "render",
            "manifest",
            "reading",
            "cache",
            "membership",
            "slice",
        ):
            row, artifacts, hashes, digest = bundle()
            if label in {"raw", "render", "manifest"}:
                name = {
                    "raw": row["raw_path"],
                    "render": "review/page-1.png",
                    "manifest": "review/manifest.json",
                }[label]
                artifacts[name] += b"tampered"
                with self.subTest(label=label), self.assertRaises(DeliverableError):
                    attach_whole_source_reviews([row], artifacts, hashes, [digest])
                continue
            cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
            extension = {
                "data": {"rows": [row]},
                "artifact_hashes": hashes,
                "whole_source_reviews": cache,
            }
            if label == "reading":
                row["reading"]["diagnostics"]["reading_order"]["status"] = "confirmed"
            elif label == "cache":
                cache[review_key(row)]["manifest"]["page_reviews"][0][
                    "body_checked"
                ] = False
            elif label == "membership":
                cache.clear()
            else:
                hashes["review/page-1.png"]["sha256"] = "0" * 64
            with self.subTest(label=label), self.assertRaises(DeliverableError):
                validate_whole_source_reviews(extension)

    def test_other_version_cannot_reuse_same_source_review(self):
        row, artifacts, hashes, digest = bundle()
        cache = attach_whole_source_reviews([row], artifacts, hashes, [digest])
        row["version_id"] = "synthetic-other-version"
        with self.assertRaises(DeliverableError):
            validate_whole_source_reviews(
                {
                    "data": {"rows": [row]},
                    "artifact_hashes": hashes,
                    "whole_source_reviews": cache,
                }
            )


if __name__ == "__main__":
    unittest.main()
