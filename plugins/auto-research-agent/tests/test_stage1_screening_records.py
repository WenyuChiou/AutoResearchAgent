"""Output-only candidate archives: real saved Core, authored HTTP, no live calls."""

import csv
import io
import sys
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "cli"))

from research_workspace.canonical_delivery import project_canonical_atlas
from research_workspace.view import write_workspace
from research_workspace_native.project_stage1_delivery import _io_root
from research_workspace_native.project_stage1_library import _project
from research_workspace_native.project_stage2 import ProjectStage2Service
from research_workspace_native.session_api import SessionApiError
from stage1_deliverable.common import (
    DeliverableError,
    canonical,
    inventory,
    private_output,
    sha,
    write_json,
)
from stage1_deliverable.package import (
    _scan_input,
    _scan_journal,
    _scan_screening_artifact,
    build,
)
from stage1_deliverable.package import validate as validate_package
from stage1_deliverable.records import CLASSIFICATION, FINDINGS, validate_records
from stage1_deliverable.screening_records import artifacts, read_artifacts
from stage1_ledger.validation import validate_run
from stage1_coverage.run import CoverageLedger
from stage2_common import Stage2Error
from stage2_workflow import build_stage2_seed
from test_stage1_core_records import core_case, produced
from test_workspace_project_intake import bound_body, compile_body, submit_body


class ScreeningRecordsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = core_case()
        cls.case, cls.fixture, _, _ = cls.context.__enter__()
        cls.addClassCleanup(cls.context.__exit__, None, None, None)
        _, cls.original_path, cls.original = produced(cls.case, cls.fixture)
        ledger = cls.fixture["ledger"]
        coverage = ledger.coverage_state()
        planned = list(coverage.queries)[2]
        query = ledger.start_planned(planned)
        backend = ledger.start(
            "backend",
            {"input": coverage.arguments(planned)},
            backend="openalex",
            parent_id=query,
        )
        cls.metadata = [
            {
                "title": "Unread " + name,
                "authors": ["Fixture Author"],
                "year": 2025,
                "doi": "10.5555/unread-" + name,
                "version": "v1",
            }
            for name in ("excluded", "pending", "unscreened")
        ]
        raw = ledger.save_bytes(
            canonical([*cls.metadata, cls.metadata[0]]), producer=backend
        )
        empty = ledger.save_bytes(b"", producer=backend)
        ledger.finish(
            backend,
            outcome="success_nonempty",
            http_status=200,
            exit_code=0,
            stdout=empty,
            stderr=empty,
            records=raw,
        )
        ledger.complete_query(query)
        ledger.extract()
        for name, status in (
            ("excluded", "include"),
            ("excluded", "exclude"),
            ("pending", "pending"),
        ):
            work = next(
                c["work_id"]
                for c in ledger.candidates().values()
                if c["discoveries"][-1]["record"]["title"] == "Unread " + name
            )
            ledger.decide(
                work,
                status,
                reason="authored-fixture",
                rationale="Saved " + status + " reason.",
                evidence_refs=[raw],
            )
        ledger.checkpoint()
        state = cls.case.intake._load(
            cls.case.intake._principal("alice-token"), cls.fixture["ref"]
        )
        run = state["stage1_runs"][state["stage1_active"]]
        report = validate_run(ledger.root)
        if not report["valid"]:
            raise AssertionError(report)
        library = _project(
            cls.fixture["ref"],
            cls.case.intake.binding(cls.fixture["ref"]),
            run,
            ledger,
            report,
        )
        cls.records = deepcopy(cls.original)
        cls.records.update(
            schema_version="1.2.0",
            screening_library=library,
            screening_artifacts=artifacts(library),
            screening_journal={
                "path": "screening_events.jsonl",
                "sha256": library["binding"]["journal_sha256"],
            },
        )
        for paper in cls.records["papers"]:
            paper["reading_status"] = "human-reviewed"
        known = {p["work_id"] for p in cls.records["papers"]}
        for row in library["rows"]:
            if row["work_id"] in known:
                continue
            metadata = row["metadata"]
            cls.records["papers"].append(
                {
                    "work_id": row["work_id"],
                    "version_id": row["selected_version_id"],
                    "title": metadata["title"],
                    "authors": metadata["authors"],
                    "year": metadata["year"],
                    "doi": metadata.get("doi"),
                    "venue": "Unknown",
                    "url": metadata.get("url"),
                    "evidence_level": "metadata",
                    "source_ids": [],
                    "claim_ids": [],
                    "roles": [],
                    "classification": dict.fromkeys(CLASSIFICATION, "Unknown"),
                    "findings": dict.fromkeys(FINDINGS, "Unknown"),
                    "reading_status": "discovery-only",
                }
            )
            first = row["discoveries"][0]
            started = ledger.event(first["query_id"])
            for decision in row["decision_history"]:
                cls.records["screening"].append(
                    {
                        "decision_id": decision["event_id"],
                        "work_id": row["work_id"],
                        "version_id": row["selected_version_id"],
                        "status": decision["new_decision"],
                        "reason": decision["rationale"],
                        "query": started["arguments"].get("query")
                        or canonical(started["arguments"]).decode(),
                        "discovery_path": first["discovery_id"],
                        "observed_at": decision["created_at"],
                    }
                )
        cls.records_path = cls.original_path.parent / "authored-screening-1.2.json"
        write_json(cls.records_path, cls.records)
        (cls.records_path.parent / "screening_events.jsonl").write_bytes(
            (ledger.root / "stage_events.jsonl").read_bytes()
        )
        for item in cls.records["screening_artifacts"]:
            target = cls.records_path.parent / item["artifact_path"]
            target.parent.mkdir(exist_ok=True)
            raw = ledger.read_ref(item["source_ref"])
            if target.exists():
                if target.read_bytes() != raw:
                    raise AssertionError("archive collision")
            else:
                target.write_bytes(raw)

    def test_all_unscreened_library_builds_empty_screening_views_without_promotion(
        self,
    ):
        case = self.case
        created = case.intake.create(
            "alice-token", {"key": "unscreened", "topic": "Authored unread studies"}
        )
        ref = created["project_ref"]
        case.intake.submit(
            "alice-token", ref, submit_body(case.intake.read("alice-token", ref))
        )
        case.intake.compile(
            "alice-token", ref, compile_body(case.intake.read("alice-token", ref))
        )
        offer = case.stage1.read("alice-token", ref)["initialization_offer"]
        case.stage1.initialize(
            "alice-token",
            ref,
            dict(
                bound_body(
                    case.intake.read("alice-token", ref), "initialize-unscreened"
                ),
                offer_ref=offer["offer_ref"],
                offer_sha256=offer["offer_sha256"],
                confirmed=True,
            ),
        )
        state = case.intake._load(case.intake._principal("alice-token"), ref)
        run = state["stage1_runs"][state["stage1_active"]]
        ledger = CoverageLedger(case.intake._path(ref, run["path"]))
        coverage = ledger.coverage_state()
        planned = next(iter(coverage.queries))
        query = ledger.start_planned(planned)
        backend = ledger.start(
            "backend",
            {"input": coverage.arguments(planned)},
            backend="openalex",
            parent_id=query,
        )
        metadata = [
            dict(
                title="Authored unread " + name,
                authors=["Unknown"],
                year=None,
                version="v1",
                doi="10.5555/authored-unread-" + name,
            )
            for name in ("first", "second")
        ]
        raw = ledger.save_bytes(canonical(metadata), producer=backend)
        empty = ledger.save_bytes(b"", producer=backend)
        ledger.finish(
            backend,
            outcome="success_nonempty",
            http_status=200,
            exit_code=0,
            stdout=empty,
            stderr=empty,
            records=raw,
        )
        ledger.complete_query(query)
        self.assertEqual(ledger.extract()["works"], 2)
        ledger.checkpoint()
        report = validate_run(ledger.root)
        self.assertTrue(report["valid"], report)
        library = _project(ref, case.intake.binding(ref), run, ledger, report)
        records = deepcopy(self.records)
        records.update(
            topic="Authored unread studies",
            papers=[],
            sources=[],
            claims=[],
            screening=[],
            core_claim_history=[],
            core_artifacts=[],
            screening_library=library,
            screening_artifacts=artifacts(library),
            screening_journal=dict(
                path="screening_events.jsonl",
                sha256=library["binding"]["journal_sha256"],
            ),
        )
        for need in records["coverage"]:
            need.update(
                work_ids=[],
                stop_decision="continue",
                reason="Unknown: unread authored fixture.",
            )
        for row in library["rows"]:
            paper_metadata = row["metadata"]
            records["papers"].append(
                dict(
                    work_id=row["work_id"],
                    version_id=row["selected_version_id"],
                    title=paper_metadata["title"],
                    authors=paper_metadata["authors"],
                    year=None,
                    doi=paper_metadata["doi"],
                    venue="Unknown",
                    url=None,
                    evidence_level="metadata",
                    source_ids=[],
                    claim_ids=[],
                    roles=[],
                    reading_status="discovery-only",
                    classification=dict.fromkeys(CLASSIFICATION, "Unknown"),
                    findings=dict.fromkeys(FINDINGS, "Unknown"),
                )
            )
        root = self.records_path.parent / "all-unscreened-detailed-input"
        root.mkdir()
        path = root / "records.json"
        write_json(path, records)
        original = path.read_bytes()
        (root / "screening_events.jsonl").write_bytes(
            (ledger.root / "stage_events.jsonl").read_bytes()
        )
        for item in records["screening_artifacts"]:
            target = root / item["artifact_path"]
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(ledger.read_ref(item["source_ref"]))
        output = _io_root(
            private_output(case.fixture_root / "all-unscreened-detailed-package")
        )
        receipt = build(path, output, sha(original))
        self.assertEqual(
            validate_package(output, receipt["manifest_sha256"])["status"], "passed"
        )
        self.assertEqual((output / "records.original.json").read_bytes(), original)
        expected_headers = list(self.records["screening"][0])
        rows = list(
            csv.reader(
                io.StringIO(
                    (output / "search_and_screening.csv").read_text("utf-8-sig")
                )
            )
        )
        self.assertEqual(rows, [expected_headers])
        workbook = load_workbook(output / "literature_catalog.xlsx", read_only=True)
        try:
            self.assertEqual(
                list(workbook["Screening"].values), [tuple(expected_headers)]
            )
            self.assertEqual(workbook["Screening library"].max_row, 3)
        finally:
            workbook.close()
        index, _ = project_canonical_atlas(
            output, "all-unscreened", receipt["manifest_sha256"]
        )
        self.assertEqual(index["screening_library"], library)
        self.assertFalse(library["stage1_complete"])
        self.assertFalse(library["execution_authorized"])
        self.assertEqual(library["counts"]["not_screened"], 2)
        self.assertEqual(library["counts"]["core"]["decisions"], 0)
        self.assertTrue(
            all(
                p["reading_status"] == "discovery-only"
                and p["findings"]["main_findings"] == "Unknown"
                and not p["source_ids"]
                and not p["claim_ids"]
                for p in index["papers"]
            )
        )
        for item in records["screening_artifacts"]:
            self.assertEqual(
                (output / item["artifact_path"]).read_bytes(),
                ledger.read_ref(item["source_ref"]),
            )
        case.runner.assert_not_called()

    def test_stage2_refuses_valid_screening_package_before_offer_or_output(self):
        output = _io_root(
            private_output(self.case.fixture_root / "screening-refusal-package")
        )
        receipt = build(self.records_path, output, sha(self.records_path.read_bytes()))
        self.assertEqual(
            validate_package(output, receipt["manifest_sha256"])["status"], "passed"
        )
        service = ProjectStage2Service(self.case.intake)
        history = self.case.intake.history("alice-token", self.fixture["ref"])
        state = self.case.intake.read("alice-token", self.fixture["ref"])
        with self.assertRaisesRegex(SessionApiError, "consumer-unconnected"):
            service.register(
                self.fixture["ref"],
                deliverable_dir=output,
                expected_manifest_sha256=receipt["manifest_sha256"],
                handoff_path=self.original_path,
                expected_handoff_sha256=sha(self.original_path.read_bytes()),
                resources_path=self.records_path,
                expected_resources_sha256=sha(self.records_path.read_bytes()),
                policy_ref=dict(
                    route="stage2-general-v3/daily_v3", execution_authorized=False
                ),
            )
        self.assertEqual(
            self.case.intake.history("alice-token", self.fixture["ref"]), history
        )
        self.assertEqual(
            self.case.intake.read("alice-token", self.fixture["ref"]), state
        )
        seed = self.case.fixture_root / "screening-refused-stage2-seed"
        with self.assertRaisesRegex(Stage2Error, "consumer-unconnected"):
            build_stage2_seed(
                output,
                receipt["manifest_sha256"],
                self.original_path,
                sha(self.original_path.read_bytes()),
                self.original_path,
                self.records_path,
                seed,
            )
        self.assertFalse(seed.exists())
        self.case.runner.assert_not_called()

    def test_all_unscreened_package_keeps_empty_screening_and_discovery_library(self):
        intake = self.case.intake
        ref = intake.create(
            "alice-token", {"key": "all-unscreened", "topic": self.original["topic"]}
        )["project_ref"]
        intake.submit("alice-token", ref, submit_body(intake.read("alice-token", ref)))
        intake.compile(
            "alice-token", ref, compile_body(intake.read("alice-token", ref))
        )
        offer = self.case.stage1.read("alice-token", ref)["initialization_offer"]
        self.case.stage1.initialize(
            "alice-token",
            ref,
            dict(
                bound_body(intake.read("alice-token", ref), "initialize"),
                offer_ref=offer["offer_ref"],
                offer_sha256=offer["offer_sha256"],
                confirmed=True,
            ),
        )
        state = intake._load(intake._principal("alice-token"), ref)
        run = state["stage1_runs"][state["stage1_active"]]
        ledger = CoverageLedger(intake._path(ref, run["path"]))
        coverage = ledger.coverage_state()
        planned = next(iter(coverage.queries))
        query = ledger.start_planned(planned)
        backend = ledger.start(
            "backend",
            {"input": coverage.arguments(planned)},
            backend="openalex",
            parent_id=query,
        )
        raw = ledger.save_bytes(canonical(self.metadata), producer=backend)
        empty = ledger.save_bytes(b"", producer=backend)
        ledger.finish(
            backend,
            outcome="success_nonempty",
            http_status=200,
            exit_code=0,
            stdout=empty,
            stderr=empty,
            records=raw,
        )
        ledger.complete_query(query)
        ledger.extract()
        ledger.checkpoint()
        report = validate_run(ledger.root)
        self.assertTrue(report["valid"], report)
        library = _project(ref, intake.binding(ref), run, ledger, report)
        records = deepcopy(self.records)
        for key in (
            "papers",
            "sources",
            "claims",
            "screening",
            "core_claim_history",
            "core_artifacts",
        ):
            records[key] = []
        records.update(
            screening_library=library,
            screening_artifacts=artifacts(library),
            screening_journal={
                "path": "screening_events.jsonl",
                "sha256": library["binding"]["journal_sha256"],
            },
        )
        for row in library["rows"]:
            metadata = row["metadata"]
            records["papers"].append(
                dict(
                    work_id=row["work_id"],
                    version_id=row["selected_version_id"],
                    title=metadata["title"],
                    authors=metadata["authors"],
                    year=metadata["year"],
                    doi=metadata.get("doi"),
                    venue="Unknown",
                    url=None,
                    evidence_level="metadata",
                    source_ids=[],
                    claim_ids=[],
                    roles=[],
                    classification=dict.fromkeys(CLASSIFICATION, "Unknown"),
                    findings=dict.fromkeys(FINDINGS, "Unknown"),
                    reading_status="discovery-only",
                )
            )
        for need in records["coverage"]:
            need["work_ids"] = []
        inputs = self.case.fixture_root / "all-unscreened-input"
        inputs.mkdir()
        path = inputs / "records.json"
        write_json(path, records)
        (inputs / "screening_events.jsonl").write_bytes(
            (ledger.root / "stage_events.jsonl").read_bytes()
        )
        for item in records["screening_artifacts"]:
            target = inputs / item["artifact_path"]
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(ledger.read_ref(item["source_ref"]))
        validate_records(records)
        self.assertEqual(records["screening"], [])
        self.assertEqual(library["counts"]["not_screened"], len(self.metadata))
        output = _io_root(
            private_output(self.case.fixture_root / "all-unscreened-package")
        )
        receipt = build(path, output, sha(path.read_bytes()))
        self.assertEqual(
            validate_package(output, receipt["manifest_sha256"])["status"], "passed"
        )
        headers = sorted(self.records["screening"][0])
        self.assertEqual(
            list(
                csv.reader(
                    io.StringIO(
                        (output / "search_and_screening.csv").read_text(
                            encoding="utf-8-sig"
                        )
                    )
                )
            ),
            [headers],
        )
        workbook = load_workbook(output / "literature_catalog.xlsx", read_only=True)
        try:
            self.assertEqual(list(workbook["Screening"].values), [tuple(headers)])
            self.assertEqual(
                workbook["Screening library"].max_row, len(self.metadata) + 1
            )
        finally:
            workbook.close()
        self.case.runner.assert_not_called()

    def test_mixed_package_atlas_and_wiki_preserve_all_candidates_without_execution(
        self,
    ):
        validate_records(self.records)
        read_artifacts(self.records, self.records_path.parent)
        output = _io_root(private_output(self.case.fixture_root / "screening-package"))
        receipt = build(self.records_path, output, sha(self.records_path.read_bytes()))
        self.assertEqual(
            validate_package(output, receipt["manifest_sha256"])["status"], "passed"
        )
        before = inventory(output)
        with (
            patch(
                "stage1_deliverable.package.build",
                side_effect=AssertionError("rebuild"),
            ),
            patch(
                "stage1_deliverable.package.validate",
                side_effect=AssertionError("replay"),
            ),
            patch(
                "stage1_deliverable.sources.validate_archive",
                side_effect=AssertionError("source replay"),
            ),
        ):
            index, files = project_canonical_atlas(
                output, "saved-screening", receipt["manifest_sha256"]
            )
        self.assertEqual(inventory(output), before)
        self.assertEqual(len(index["papers"]), 5)
        self.assertEqual(len(index["sources"]), 2)
        self.assertEqual(index["screening_library"], self.records["screening_library"])
        self.assertEqual(len(index["bibliography"]["entries"]), 5)
        self.assertEqual(index["stage2_canonical_consumer"]["status"], "unconnected")
        unread = [p for p in index["papers"] if p["reading_status"] == "discovery-only"]
        self.assertEqual(len(unread), 3)
        self.assertTrue(
            all(
                p["url"] is None
                and not p["source_ids"]
                and p["findings"]["main_findings"] == "Unknown"
                for p in unread
            )
        )
        library = index["screening_library"]
        self.assertEqual(library["counts"]["not_screened"], 1)
        excluded = next(
            r for r in library["rows"] if r["metadata"]["title"] == "Unread excluded"
        )
        self.assertEqual(
            (excluded["counts"]["discoveries"], excluded["counts"]["reversals"]), (2, 1)
        )
        self.assertIn(
            b"Saved exclude reason.",
            files[
                "wiki/"
                + sha(canonical([excluded["work_id"], excluded["selected_version_id"]]))
                + ".md"
            ],
        )
        self.assertEqual(index["native_session"]["status"], "not-connected")
        self.assertFalse(index["readiness"]["stage2_execution_authorized"])
        reference = (
            Path(__file__).resolve().parents[1] / "references/research-workspace"
        )
        with self.assertRaisesRegex(
            DeliverableError, "Stage 2 delivery, bridge and receipt hash are required"
        ):
            write_workspace(
                index,
                reference,
                self.case.fixture_root / "refused-stage2",
                stage2_delivery="unused",
            )

    def test_old_guards_remain_and_discovery_cannot_be_promoted_or_hide_source(self):
        validate_records(self.original)
        for field, value in (("url", None), ("source_ids", [])):
            old = deepcopy(self.original)
            old["papers"][0][field] = value
            with self.assertRaises(DeliverableError):
                validate_records(old)
        for mutate in (
            lambda r: r["papers"][-1]["findings"].update(
                main_findings="Invented improvement"
            ),
            lambda r: r["papers"][-1].update(reading_status="human-reviewed"),
            lambda r: r["papers"][0].update(
                reading_status="discovery-only",
                source_ids=[],
                claim_ids=[],
                roles=[],
                evidence_level="metadata",
                findings=dict.fromkeys(FINDINGS, "Unknown"),
                classification=dict.fromkeys(CLASSIFICATION, "Unknown"),
            ),
        ):
            changed = deepcopy(self.records)
            mutate(changed)
            with self.assertRaises(DeliverableError):
                validate_records(changed)

    def test_history_counts_and_candidate_dropping_are_refused(self):
        for mutate in (
            lambda r: r["papers"].pop(),
            lambda r: r["screening"].pop(),
            lambda r: r["screening_library"]["rows"][-1]["counts"].update(
                discoveries=0
            ),
            lambda r: r["screening_library"]["counts"].update(not_screened=0),
            lambda r: r["screening_artifacts"].pop(),
        ):
            changed = deepcopy(self.records)
            mutate(changed)
            with self.assertRaises(DeliverableError):
                validate_records(changed)

    def test_raw_discovery_archive_tamper_is_refused(self):
        row = next(
            r
            for r in self.records["screening_library"]["rows"]
            if r["metadata"]["title"] == "Unread excluded"
        )
        ref = row["discoveries"][0]["source_ref"]
        item = next(
            a for a in self.records["screening_artifacts"] if a["source_ref"] == ref
        )

        def read(path):
            raw = (self.records_path.parent / path).read_bytes()
            return raw + b"tamper" if path == item["artifact_path"] else raw

        from stage1_deliverable.screening_records import validate_artifacts

        with self.assertRaisesRegex(DeliverableError, "raw artifact hash differs"):
            validate_artifacts(self.records, read)

    def test_core_null_scan_copy_preserves_bytes_and_still_rejects_credentials(self):
        original = canonical(self.records)
        _scan_input(self.records, original)
        journal = (self.records_path.parent / "screening_events.jsonl").read_bytes()
        _scan_journal(self.records, journal)
        self.assertEqual(
            (self.records_path.parent / "screening_events.jsonl").read_bytes(), journal
        )
        self.assertEqual(canonical(self.records), original)
        secret = deepcopy(self.records)
        secret["papers"][0]["findings"]["limitations"] = (
            'authorization="SYNTHETIC_DO_NOT_USE"'
        )
        with self.assertRaisesRegex(DeliverableError, "credential-like"):
            _scan_input(secret, canonical(secret))
        literal = deepcopy(self.records)
        literal["screening_library"]["snapshots"][2]["rows"][0]["authorization"] = (
            "null"
        )
        with self.assertRaisesRegex(DeliverableError, "credential-like"):
            _scan_input(literal, canonical(literal))
        snapshot = self.records["screening_library"]["snapshots"][2]
        ref = snapshot["persisted_refs"][0]
        artifact = next(
            a for a in self.records["screening_artifacts"] if a["source_ref"] == ref
        )
        raw = (self.records_path.parent / artifact["artifact_path"]).read_bytes()
        _scan_screening_artifact(self.records, artifact, raw)
        self.assertEqual(
            (self.records_path.parent / artifact["artifact_path"]).read_bytes(), raw
        )
        arbitrary = next(
            a for a in self.records["screening_artifacts"] if a["source_ref"] != ref
        )
        for body in (
            b'authorization="SYNTHETIC_DO_NOT_USE"',
            b'{"authorization":null}',
        ):
            with self.assertRaisesRegex(DeliverableError, "credential-like"):
                _scan_screening_artifact(self.records, arbitrary, body)
        for field, value in (("doi", "10.5555/invented"), ("venue", "Invented venue")):
            changed = deepcopy(self.records)
            changed["papers"][-1][field] = value
            with self.assertRaisesRegex(DeliverableError, "DOI or venue differs"):
                validate_records(changed)
        from stage1_deliverable.screening_records import journal_rows

        with self.assertRaisesRegex(DeliverableError, "journal hash differs"):
            journal_rows(self.records, journal + b"tamper\n")
        changed = deepcopy(self.records)
        changed["screening_library"]["rows"][0]["source_attempts"] = []
        with self.assertRaisesRegex(
            DeliverableError, "journal source/import/review history differs"
        ):
            journal_rows(changed, journal)


if __name__ == "__main__":
    unittest.main()
