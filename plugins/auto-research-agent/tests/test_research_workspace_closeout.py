"""Lossless private closeout views and source-state preservation regressions."""

import io
import json
import unittest
from copy import deepcopy
from pathlib import Path

import test_research_workspace_repair as repair_fixture
import test_research_workspace_view as view_fixture
from openpyxl import load_workbook
from research_workspace.closeout import closeout_files
from research_workspace.projection import validate_index
from research_workspace.view import render_view, write_workspace
from research_workspace.wiki import wiki_files
from stage1_deliverable.common import DeliverableError, canonical, sha
from stage1_deliverable.views import bibtex


class CloseoutTests(unittest.TestCase):
    def setUp(self):
        self.repair = repair_fixture.RepairProjectionTests()
        self.repair.setUp()
        self.addCleanup(self.repair.doCleanups)
        full = view_fixture.fixture_index()
        provenance = {**full["provenance"], **self.repair.index["provenance"]}
        paper = deepcopy(full["papers"][0])
        paper.update(self.repair.index["papers"][0])
        paper["source_ids"] = ["src-neutral"]
        paper["claim_ids"] = ["neutral-parent"]
        full.update(self.repair.index)
        full["provenance"] = provenance
        full["papers"] = [paper]
        full["edges"] = []
        full["bibliography"] = {
            "producer": "stage1_deliverable.views.bibtex",
            "records_sha256": full["provenance"]["records_sha256"],
            "all_bibtex": bibtex(full).decode(),
            "entries": [
                {
                    "work_id": paper["work_id"],
                    "version_id": paper["version_id"],
                    "bibtex": bibtex(full).decode(),
                }
            ],
        }
        self.repair.index = full
        self.index = self.repair.apply()
        validate_index(self.index)

    def test_saved_canonical_index_rebuilds_identical_exports(self):
        saved = json.loads(canonical(self.index))
        before = closeout_files(self.index)
        after = closeout_files(saved)
        self.assertEqual(set(before), set(after))
        for path in before:
            with self.subTest(path=path):
                self.assertEqual(before[path], after[path])

    def test_generic_summary_cannot_invent_case_or_owner_state(self):
        text = closeout_files(self.index)["closeout/README.zh-TW.md"].decode()
        for value in ("Beckman", "EconAgent", "SUPERD6", "Core", "排程維持暫停"):
            with self.subTest(value=value):
                self.assertNotIn(value, text)

    def test_deterministic_cross_format_identity_and_original_unknown(self):
        before = deepcopy(self.index)
        files = closeout_files(self.index)
        self.assertEqual(files, closeout_files(self.index))
        self.assertEqual(before, self.index)
        book = load_workbook(io.BytesIO(files["closeout/literature-catalog.xlsx"]))
        for name in ("Papers", "Claims", "AtomicRevisions", "CoreFindings"):
            self.assertEqual(book[name].max_row, 2)
            keys = [cell.value for cell in book[name][1]]
            rows = dict(zip(keys, [cell.value for cell in book[name][2]]))
            self.assertEqual(rows["work_id"], "neutral-work")
            self.assertEqual(rows["version_id"], "version-one")
        claims = dict(
            zip(
                [x.value for x in book["Claims"][1]],
                [x.value for x in book["Claims"][2]],
            )
        )
        self.assertEqual(claims["relation"], "unknown")
        atoms = json.loads(files["closeout/claim-revision-map.json"])
        self.assertEqual(atoms, self.index["supplement"]["data"]["atomic_revisions"])
        notes = wiki_files(self.index)
        path = "wiki/" + sha(canonical(["neutral-work", "version-one"])) + ".md"
        self.assertIn(b"neutral-atom", notes[path])
        self.assertIn(b"neutral-parent", notes[path])
        self.assertIn(b"still unknown", notes[path])
        binding = json.loads(files["closeout/manifest.json"])
        self.assertEqual(
            binding["runtime"]["original_runtime_replay"],
            "not-reexecuted-by-this-exporter",
        )
        self.assertFalse(binding["official_stage2_import_eligible"])
        self.assertEqual(binding["original_compound_counts"]["unknown"], 1)

    def test_view_binds_all_outputs_without_execution_and_refuses_tampering(self):
        root = Path(self.repair.temporary.name).resolve()
        reference = Path(__file__).parents[1] / "references/research-workspace"
        first = write_workspace(self.index, reference, root / "view-one")
        second = write_workspace(self.index, reference, root / "view-two")
        self.assertEqual(first, second)
        rebuilt = render_view(
            root / "view-one/workspace-index.json",
            root / "view-from-saved-index",
            reference,
            first["index_sha256"],
        )
        self.assertEqual(first, rebuilt)
        self.assertEqual(first["research_execution"], "not-performed")
        self.assertFalse(first["repair_binding"]["original_replay_upgraded"])
        for name, digest in first["files"].items():
            self.assertEqual(sha((root / "view-one" / name).read_bytes()), digest)
        self.assertEqual(
            (root / "view-one/references.bib").read_text(encoding="utf-8"),
            self.index["bibliography"]["all_bibtex"],
        )
        html = (root / "view-one/index.html").read_text(encoding="utf-8")
        self.assertIn("connect-src 'none'", html)
        self.assertIn('href="./workspace-closeout.css"', html)
        self.assertIn('<body class="stage1-closeout">', html)
        self.assertIn("workspace-closeout.css", first["files"])
        self.assertLess(
            html.index("workspace-repairs.js"), html.index("workspace-records.js")
        )
        payload_text = (root / "view-one/workspace-data.js").read_text(encoding="utf-8")
        payload = json.loads(
            payload_text.removeprefix("window.WORKSPACE_VIEW = ")
            .strip()
            .removesuffix(";")
        )
        for row in payload["note_paths"]:
            self.assertEqual(
                row["text"],
                (root / "view-one" / row["path"]).read_text(encoding="utf-8"),
            )
        for key, value in (
            ("official_stage2_import_eligible", True),
            (
                "source_reading",
                {"abstract": "success", "required_full_methods": "success"},
            ),
        ):
            with self.subTest(key=key):
                changed = deepcopy(self.index)
                changed["supplement"]["data"][key] = value
                with self.assertRaises(DeliverableError):
                    write_workspace(changed, reference, root / ("bad-" + key))
                self.assertFalse((root / ("bad-" + key)).exists())

    def test_absent_supplement_retains_v1_behavior(self):
        index = view_fixture.fixture_index()
        self.assertEqual(closeout_files(index), {})
        root = Path(self.repair.temporary.name).resolve()
        reference = Path(__file__).parents[1] / "references/research-workspace"
        receipt = write_workspace(index, reference, root / "legacy")
        self.assertNotIn("repair_binding", receipt)
        self.assertNotIn("workspace-repairs.js", receipt["files"])
        self.assertNotIn("workspace-closeout.css", receipt["files"])
        self.assertNotIn(
            "workspace-closeout.css",
            (root / "legacy/index.html").read_text(encoding="utf-8"),
        )
        self.assertNotIn(
            'class="stage1-closeout"',
            (root / "legacy/index.html").read_text(encoding="utf-8"),
        )
        self.assertEqual(
            index["supplement"], {"status": "not-provided", "top3_status": "pending"}
        )


if __name__ == "__main__":
    unittest.main()
