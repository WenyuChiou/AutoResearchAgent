"""Synthetic formal-selection policy and byte-export regressions."""

import csv
import io
import json
import re
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parents[1] / "cli"))

import test_research_workspace_closeout as closeout_fixture
import test_research_workspace_source_rerun_exports as rerun_fixture
from research_workspace.literature_selection import (
    derive_literature_selection,
    selection_files,
)
from research_workspace.source_rerun import attach_rerun
from research_workspace.view import write_workspace
from stage1_deliverable.common import DeliverableError, canonical, sha


def attached_index(change=None, base_change=None):
    index = rerun_fixture.fixture_index()
    if base_change:
        base_change(index)
    manifest = rerun_fixture.fixture_manifest(index)
    complete_document(manifest["data"]["rows"][0])
    if change:
        change(manifest)
    temporary = tempfile.TemporaryDirectory(prefix="literature-selection-")
    root = Path(temporary.name).resolve()
    expected = rerun_fixture.write_manifest(root, manifest)
    return temporary, attach_rerun(index, root, expected)


def complete_document(row):
    """Supply an affirmative full-body receipt for the known synthetic text."""
    reading = row["reading"]
    reading["locators"] = [
        {"start": 0, "end": reading["characters"], "type": "text", "value": "body"}
    ]
    reading["diagnostics"]["document_extent"] = {
        "scope": "full-body",
        "complete": True,
        "raw_sha256": row["raw_sha256"],
        "text_sha256": reading["text_sha256"],
        "expected_characters": reading["characters"],
        "readable_characters": reading["characters"],
    }


def change_second(status, level, identity="consistent"):
    def change(manifest):
        row = manifest["data"]["rows"][1]
        row["reading"]["status"] = status
        row["reading"]["evidence_level"] = level
        row["reading"]["identity_status"] = identity
        if status not in {"extracted", "identity-mismatch"}:
            manifest["files"].pop(row.pop("extracted_path"))
            row["reading"].update(
                characters=0,
                text_sha256=None,
                locators=[],
                error={"type": "SyntheticError", "message": status},
            )
        else:
            row["reading"]["locators"][0]["value"] = "body"
            if level == "full-text":
                complete_document(row)
            if status == "identity-mismatch":
                row["reading"]["observed_identity"] = {
                    "title": "different work",
                    "doi": "",
                }

    return change


class LiteratureSelectionTests(unittest.TestCase):
    def test_full_text_without_affirmative_extent_stays_pending(self):
        def missing_extent(manifest):
            manifest["data"]["rows"][0]["reading"]["diagnostics"] = {}

        temporary, index = attached_index(missing_extent)
        self.addCleanup(temporary.cleanup)
        before = canonical(index)
        result = derive_literature_selection(index)
        row = result["rows"][0]
        self.assertEqual(row["status"], "pending")
        self.assertEqual(row["eligible_source_ids"], [])
        self.assertEqual(
            row["source_binding_references"][0]["body_completeness"]["status"],
            "pending",
        )
        self.assertEqual(result["rule_version"], "1.1.0")
        self.assertEqual(canonical(index), before)
        self.assertEqual(index["claims"][0]["support"], "Unknown")

    def test_inconsistent_body_extent_is_excluded_without_rewriting_evidence(self):
        def inconsistent_extent(manifest):
            extent = manifest["data"]["rows"][0]["reading"]["diagnostics"][
                "document_extent"
            ]
            extent["expected_characters"] += 1

        temporary, index = attached_index(inconsistent_extent)
        self.addCleanup(temporary.cleanup)
        before = canonical(index)
        row = derive_literature_selection(index)["rows"][0]
        self.assertEqual(row["status"], "excluded")
        self.assertEqual(
            row["source_binding_references"][0]["body_completeness"]["status"],
            "incomplete",
        )
        self.assertEqual(canonical(index), before)

    def test_empty_legacy_index_exports_headers_without_inventing_records(self):
        index = rerun_fixture.view_fixture.fixture_index()
        index.update(papers=[], sources=[], claims=[])
        index["bibliography"].update(
            entries=[], all_bibtex=rerun_fixture.bibtex(index).decode()
        )
        exported = selection_files(index)
        selection = json.loads(exported["literature/selection.json"])
        self.assertEqual(
            selection["counts"],
            {"screened": 0, "included": 0, "excluded": 0, "pending": 0},
        )
        self.assertEqual(exported["literature/included.bib"], b"")
        self.assertEqual(exported["literature/screening.bib"], b"")
        book = load_workbook(
            io.BytesIO(exported["literature/catalog.xlsx"]), read_only=True
        )
        self.assertEqual(
            {sheet.title: sheet.max_row for sheet in book},
            {"Included": 1, "Screening": 1, "Selection": 1},
        )
        csv_rows = list(
            csv.reader(
                io.StringIO(exported["literature/selection.csv"].decode("utf-8-sig"))
            )
        )
        self.assertEqual(len(csv_rows), 1)
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder).resolve() / "empty-view"
            reference = Path(__file__).parents[1] / "references/research-workspace"
            receipt = write_workspace(index, reference, output)
            self.assertIn("literature/selection.json", receipt["files"])

    def test_full_text_is_included_despite_unknown_claim(self):
        temporary, index = attached_index()
        self.addCleanup(temporary.cleanup)
        before = canonical(index)
        result = derive_literature_selection(index)
        self.assertEqual(
            result["counts"],
            {"screened": 2, "included": 1, "excluded": 1, "pending": 0},
        )
        self.assertEqual(
            result["included_identities"],
            [{"work_id": "fixture-0", "version_id": "version-0"}],
        )
        self.assertEqual(index["claims"][0]["support"], "Unknown")
        self.assertEqual(before, canonical(index))
        self.assertFalse(result["claim_assessments_changed"])
        self.assertFalse(result["coverage_complete"])
        self.assertFalse(result["official_stage2_import_eligible"])
        self.assertFalse(result["research_execution"])

    def test_two_versions_have_distinct_bibliography_keys_without_record_changes(self):
        def two_versions(index):
            index["papers"][1]["work_id"] = "fixture-0"
            index["sources"][1]["work_id"] = "fixture-0"
            rerun_fixture.refresh_bibliography(index)

        temporary, index = attached_index(
            change_second("extracted", "full-text"), two_versions
        )
        self.addCleanup(temporary.cleanup)
        before = canonical(index)
        exported = selection_files(index)
        selection = json.loads(exported["literature/selection.json"])
        keys = {row["citation_key"] for row in selection["rows"]}
        self.assertEqual(len(keys), 2)
        for name in ("included.bib", "screening.bib"):
            actual = re.findall(
                r"@misc\{([^,]+),", exported["literature/" + name].decode()
            )
            self.assertEqual(set(actual), keys)
            self.assertEqual(len(actual), 2)
        self.assertEqual(
            {(row["work_id"], row["version_id"]) for row in selection["rows"]},
            {("fixture-0", "version-0"), ("fixture-0", "version-1")},
        )
        self.assertEqual(canonical(index), before)

    def test_nonqualifying_current_states_have_specific_dispositions(self):
        cases = (
            (change_second("extracted", "abstract"), "excluded", "abstract-only"),
            (change_second("extracted", "metadata"), "excluded", "metadata-only"),
            (change_second("inaccessible", "metadata"), "excluded", "inaccessible"),
            (
                change_second("inaccessible", "metadata", "unverified"),
                "excluded",
                "inaccessible",
            ),
            (
                change_second("failed-engineering", "metadata"),
                "excluded",
                "broken-parser",
            ),
            (
                change_second("extracted", "full-text", "unverified"),
                "pending",
                "identity-unverified",
            ),
            (
                change_second("identity-mismatch", "full-text", "mismatch"),
                "excluded",
                "identity-mismatch",
            ),
        )
        for change, status, reason in cases:
            with self.subTest(reason=reason):
                temporary, index = attached_index(change)
                try:
                    row = derive_literature_selection(index)["rows"][1]
                    self.assertEqual(row["status"], status)
                    self.assertIn(reason, row["reasons"])
                    self.assertEqual(row["eligible_source_ids"], [])
                finally:
                    temporary.cleanup()

    def test_omission_diagnostics_block_admission_and_remain_source_bound(self):
        for diagnostics, reason in (
            ({"omitted_pages": [3, 4]}, "source-pages-omitted"),
            ({"incomplete": "parser stopped early"}, "source-incomplete"),
        ):
            with self.subTest(reason=reason):

                def diagnose(manifest):
                    manifest["data"]["rows"][0]["reading"]["diagnostics"] = diagnostics

                temporary, index = attached_index(diagnose)
                try:
                    row = derive_literature_selection(index)["rows"][0]
                    self.assertEqual(row["status"], "excluded")
                    self.assertIn(reason, row["reasons"])
                    binding = row["source_binding_references"][0]
                    self.assertEqual(binding["diagnostics"], diagnostics)
                    self.assertIn(reason, binding["failure_reasons"])
                finally:
                    temporary.cleanup()

    def test_rerun_source_outside_canonical_paper_membership_is_rejected(self):
        def wrong_membership(index):
            index["papers"][0]["source_ids"] = ["source-1"]

        temporary, index = attached_index(base_change=wrong_membership)
        self.addCleanup(temporary.cleanup)
        with self.assertRaisesRegex(DeliverableError, "canonical paper membership"):
            derive_literature_selection(index)

    def test_abstract_cannot_be_promoted_and_hash_or_version_tampering_rejects(self):
        temporary, index = attached_index()
        self.addCleanup(temporary.cleanup)
        mutations = (
            lambda value: value["source_rerun"]["data"]["rows"][1]["reading"].update(
                evidence_level="full-text"
            ),
            lambda value: value["source_rerun"]["data"]["rows"][0].update(
                raw_sha256="a" * 64
            ),
            lambda value: value["source_rerun"]["data"]["rows"][0]["reading"].update(
                text_sha256="b" * 64
            ),
            lambda value: value["source_rerun"]["data"]["rows"][0].update(
                version_id="different-version"
            ),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                changed = deepcopy(index)
                mutate(changed)
                with self.assertRaises(DeliverableError):
                    derive_literature_selection(changed)

    def test_schema_one_and_two_remain_pending_source_review(self):
        v1 = rerun_fixture.fixture_index()
        first = derive_literature_selection(v1)
        self.assertEqual(first["counts"]["pending"], len(v1["papers"]))
        self.assertTrue(
            all(row["reasons"] == ["pending-source-review"] for row in first["rows"])
        )
        pending_book = load_workbook(
            io.BytesIO(selection_files(v1)["literature/catalog.xlsx"]), read_only=True
        )
        self.assertEqual(pending_book["Included"].max_row, 1)

        closeout = closeout_fixture.CloseoutTests()
        closeout.setUp()
        self.addCleanup(closeout.doCleanups)
        second = derive_literature_selection(closeout.index)
        self.assertEqual(second["counts"]["pending"], len(closeout.index["papers"]))
        self.assertTrue(all(row["status"] == "pending" for row in second["rows"]))

    def test_view_payload_and_exports_bind_one_selection_without_execution(self):
        temporary, index = attached_index()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name).resolve()
        output = root / "view"
        reference = Path(__file__).parents[1] / "references/research-workspace"
        receipt = write_workspace(index, reference, output, source_rerun_root=root)
        data = (output / "workspace-data.js").read_text(encoding="utf-8")
        payload = json.loads(
            data.removeprefix("window.WORKSPACE_VIEW = ").strip().removesuffix(";")
        )
        selection = json.loads((output / "literature/selection.json").read_bytes())
        self.assertEqual(payload["literature_selection"], selection)
        self.assertEqual(canonical(payload["index"]), canonical(index))
        self.assertEqual(receipt["research_execution"], "not-performed")
        for name, raw in selection_files(index).items():
            self.assertEqual(receipt["files"][name], sha(raw))
            self.assertEqual((output / name).read_bytes(), raw)
        module = (
            Path(__file__).parents[1] / "cli/research_workspace/literature_selection.py"
        )
        self.assertEqual(
            receipt["adapter_sources"]["literature_selection.py"],
            sha(module.read_bytes()),
        )
        self.assertEqual(
            receipt["adapter_sources"]["body_completeness.py"],
            sha(module.with_name("body_completeness.py").read_bytes()),
        )

    def test_exports_are_deterministic_and_reconcile_exact_identities(self):
        temporary, index = attached_index()
        self.addCleanup(temporary.cleanup)
        before = canonical(index)
        first = selection_files(index)
        self.assertEqual(first, selection_files(index))
        self.assertEqual(first, selection_files(json.loads(canonical(index))))
        self.assertEqual(before, canonical(index))
        self.assertEqual(
            set(first),
            {
                "literature/selection.json",
                "literature/selection.csv",
                "literature/included.bib",
                "literature/screening.bib",
                "literature/selection.md",
                "literature/catalog.xlsx",
            },
        )
        selection = json.loads(first["literature/selection.json"])
        expected_all = {
            (row["work_id"], row["version_id"]) for row in selection["rows"]
        }
        expected_included = {
            (row["work_id"], row["version_id"])
            for row in selection["rows"]
            if row["status"] == "included"
        }
        csv_rows = list(
            csv.DictReader(
                io.StringIO(first["literature/selection.csv"].decode("utf-8-sig"))
            )
        )
        self.assertEqual(
            {(row["work_id"], row["version_id"]) for row in csv_rows}, expected_all
        )
        workbook = load_workbook(
            io.BytesIO(first["literature/catalog.xlsx"]), read_only=True
        )
        for sheet, expected in (
            ("Included", expected_included),
            ("Screening", expected_all),
        ):
            values = list(workbook[sheet].iter_rows(values_only=True))
            headers = list(values[0])
            actual = {
                (row[headers.index("work_id")], row[headers.index("version_id")])
                for row in values[1:]
            }
            self.assertEqual(actual, expected)
        included_bib = first["literature/included.bib"].decode()
        screening_bib = first["literature/screening.bib"].decode()
        self.assertEqual(included_bib.count("@misc{"), len(expected_included))
        self.assertEqual(screening_bib.count("@misc{"), len(expected_all))
        self.assertEqual(
            set(re.findall(r"@misc\{([^,]+),", screening_bib)),
            {row["citation_key"] for row in selection["rows"]},
        )
        self.assertIn("Version version-0", screening_bib)
        self.assertIn("Version version-1", screening_bib)


if __name__ == "__main__":
    unittest.main()
