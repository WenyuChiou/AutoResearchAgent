"""Offline reader regressions for rerunning complete saved-source archives."""

import json
import sys
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "cli"))

import test_research_workspace_view as view_fixture
from research_workspace.source_rerun import (
    _saved_public_import,
    attach_rerun,
    build_rerun,
    rerun_files,
    validate_rerun_index,
)
from research_hub.source_fetch_saved import import_saved_public_source
from stage1_deliverable.common import DeliverableError, canonical, sha
from stage1_deliverable.sources import artifact_map, receipt_digest


PARSER = b"""from dataclasses import dataclass

@dataclass
class Item:
    text: str
    evidence_level: str
    observed_title: str
    observed_doi: str
    locators: list
    diagnostics: dict
    bibliographic_metadata: dict

def _extract_html(raw, url):
    text = "Evidence saved full text"
    return Item(text, "full-text", "Evidence", "10.1000/synthetic",
                [{"start": 0, "end": len(text), "type": "text", "value": "body"}],
                {"parser": "synthetic-local"},
                {"title": "Evidence", "doi": "10.1000/synthetic", "volume": "7"})

def _extract_pdf(raw):
    raise RuntimeError("PDF is outside this synthetic fixture")

def _identity(expected_doi, expected_title, observed_doi, observed_title):
    return "consistent" if expected_title == observed_title else "mismatch"
"""


def receipt(source_id, raw, *, status, http_status, attempts=1):
    url = f"https://example.test/{source_id}"
    rows = []
    if attempts == 2:
        rows.append(
            {
                "sequence": 1,
                "purpose": "synthetic redirect",
                "url": url,
                "final_url": url + "/moved",
                "requested_at": "2026-10-06T00:00:00Z",
                "received_at": "2026-10-06T00:00:01Z",
                "http_status": 302,
                "content_type": "text/html",
                "outcome": "redirect",
                "response_bytes": 0,
                "response_truncated": False,
                "raw_path": None,
                "raw_sha256": None,
                "error": None,
            }
        )
    rows.append(
        {
            "sequence": len(rows) + 1,
            "purpose": "synthetic saved response",
            "url": url,
            "final_url": url + "/final",
            "requested_at": "2026-10-06T00:00:02Z",
            "received_at": "2026-10-06T00:00:03Z",
            "http_status": http_status,
            "content_type": "text/html",
            "outcome": status,
            "response_bytes": len(raw),
            "response_truncated": False,
            "raw_path": "saved.html",
            "raw_sha256": sha(raw),
            "error": None,
        }
    )
    value = {
        "schema_version": "source-fetch-result/v1",
        "request": {
            "operation": "fetch",
            "doi": None,
            "url": url,
            "title": source_id,
            "output_dir": f"/synthetic/{source_id}",
            "public_only": True,
        },
        "receipt_sha256": "",
        "status": status,
        "evidence_level": "metadata",
        "source_url": url,
        "final_url": url + "/final",
        "retrieved_at": "2026-10-06T00:00:03Z",
        "expected_identity": {"doi": None, "title": source_id},
        "observed_identity": {"doi": None, "title": None},
        "identity_status": "unverified",
        "source_version": "sha256:" + sha(raw),
        "attempts": rows,
        "raw_path": "saved.html",
        "raw_sha256": sha(raw),
        "extracted_text_path": None,
        "extracted_text_sha256": None,
        "locators": [],
        "errors": [],
        "output_dir": f"/synthetic/{source_id}",
    }
    value["receipt_sha256"] = receipt_digest(value)
    return value


class SavedReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="saved-reader-", dir=Path(tempfile.gettempdir()).resolve()
        )
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.package = self.root / "package"
        self.package.mkdir()
        self.parser = self.root / "parser.py"
        self.parser.write_bytes(PARSER)
        self.parser_sha = sha(PARSER)
        self.raw = {
            "source-0": b"<html><body>successful saved response</body></html>",
            "source-1": b"<html><body>forbidden saved response</body></html>",
        }
        self.index = view_fixture.fixture_index()
        receipts = (
            receipt(
                "source-0",
                self.raw["source-0"],
                status="available",
                http_status=200,
                attempts=2,
            ),
            receipt(
                "source-1", self.raw["source-1"], status="paywalled", http_status=403
            ),
        )
        self.index["sources"] = []
        for number, saved in enumerate(receipts):
            archive = self.package / "deliverable/sources" / f"source-{number}"
            archive.mkdir(parents=True)
            (archive / "original.json").write_bytes(canonical(saved))
            (archive / "saved.html").write_bytes(self.raw[f"source-{number}"])
            self.index["sources"].append(
                {
                    "source_id": f"source-{number}",
                    "work_id": f"fixture-{number}",
                    "version_id": f"version-{number}",
                    "receipt": saved,
                    "result_sha256": sha(canonical(saved)),
                }
            )

    def build(self, name="rerun"):
        output = self.root / name
        digest = build_rerun(
            self.index, self.package, output, self.parser, self.parser_sha
        )
        manifest = json.loads((output / "source-rerun-manifest.json").read_bytes())
        return output, digest, manifest

    def test_success_and_non200_preserve_saved_attempt_history(self):
        output, digest, manifest = self.build()
        attach_rerun(self.index, output, digest)
        self.assertEqual(digest, sha(canonical(manifest)))
        success, inaccessible = manifest["data"]["rows"]
        self.assertEqual(
            success["original_attempts"],
            self.index["sources"][0]["receipt"]["attempts"],
        )
        self.assertEqual(success["reading"]["status"], "extracted")
        self.assertEqual(success["metadata"]["doi"], "10.1000/synthetic")
        self.assertEqual(success["metadata"]["volume"], "7")
        self.assertEqual(inaccessible["reading"]["status"], "inaccessible")
        self.assertEqual(inaccessible["reading"]["error"]["type"], "PermissionError")

    def test_parser_hash_is_pinned_before_reading_sources(self):
        with self.assertRaisesRegex(DeliverableError, "parser hash differs"):
            build_rerun(
                self.index,
                self.package,
                self.root / "wrong-parser",
                self.parser,
                "0" * 64,
            )

    def test_changed_saved_source_bytes_are_rejected(self):
        path = self.package / "deliverable/sources/source-0/saved.html"
        path.write_bytes(b"changed after receipt")
        with self.assertRaisesRegex(DeliverableError, "raw source hash differs"):
            build_rerun(
                self.index,
                self.package,
                self.root / "changed-source",
                self.parser,
                self.parser_sha,
            )


def _bind_history(reader):
    source = reader.index["sources"][0]
    saved = source["receipt"]
    saved["receipt_sha256"] = receipt_digest(saved)
    original = canonical(saved)
    (reader.package / "deliverable/sources/source-0/original.json").write_bytes(
        original
    )
    source["result_sha256"] = sha(original)


def _offline_reader(test_case):
    reader = SavedReaderTests()
    reader.setUp()
    test_case.addCleanup(reader.doCleanups)
    source = reader.index["sources"][0]
    acquisition = source["receipt"]["attempts"][-1]
    offline = deepcopy(acquisition)
    offline.update(
        sequence=acquisition["sequence"] + 1,
        purpose="offline-parser-replay",
        http_status=None,
        outcome="parsed",
        raw_path="replay.html",
    )
    source["receipt"]["attempts"].append(offline)
    (reader.package / "deliverable/sources/source-0/replay.html").write_bytes(
        reader.raw["source-0"]
    )
    _bind_history(reader)
    return reader


def _saved_import_reader(test_case):
    reader = SavedReaderTests()
    reader.setUp()
    test_case.addCleanup(reader.doCleanups)
    input_root = reader.root / "saved-input"
    input_root.mkdir()
    raw = (
        b"<html><head><title>Evidence</title></head><body><article><h1>Evidence</h1>"
        + (b"<p>Saved public scholarly evidence with reproducible details.</p>" * 40)
        + b"</article></body></html>"
    )
    (input_root / "source.html").write_bytes(raw)
    manifest = {
        "schema_version": "saved-public-source-input/v1",
        "url": "https://example.test/source-0",
        "final_url": "https://example.test/source-0/final",
        "doi": "10.1000/synthetic",
        "title": "Evidence",
        "raw_path": "source.html",
        "raw_sha256": sha(raw),
        "content_type": "text/html",
        "original_acquired_at": "2026-10-06T00:00:00Z",
    }
    manifest_bytes = canonical(manifest)
    manifest_path = input_root / "manifest.json"
    manifest_path.write_bytes(manifest_bytes)
    imported_root = reader.root / "sdk-import"
    imported = import_saved_public_source(
        input_manifest=manifest_path,
        expected_manifest_sha256=sha(manifest_bytes),
        output_dir=imported_root,
    ).to_dict()
    archive = reader.package / "deliverable/sources/source-0"
    for source_path, relative in artifact_map(imported).items():
        target = archive / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(Path(source_path).read_bytes())
    original = (imported_root / "source-fetch-result.json").read_bytes()
    (archive / "original.json").write_bytes(original)
    source = reader.index["sources"][0]
    source["receipt"] = imported
    source["result_sha256"] = sha(original)
    return reader


def _legacy_saved_import_rerun(reader):
    """Complete fixture emitted by the edec9ddd writer, before saved-import replay."""
    output = reader.root / "source-rerun"
    rows, files = [], {}
    base_sha = sha(canonical(reader.index))
    for source in reader.index["sources"]:
        saved = source["receipt"]
        selected = saved["attempts"][-1]
        archive = reader.package / "deliverable/sources" / source["source_id"]
        raw = (archive / artifact_map(saved)[selected["raw_path"]]).read_bytes()
        raw_name = "sources/" + source["source_id"] + "/raw.html"
        files[raw_name] = raw
        row = {key: source[key] for key in ("work_id", "version_id", "source_id")}
        row.update(
            original_result_sha256=source["result_sha256"],
            previous_status=saved["status"],
            previous_evidence_level=saved["evidence_level"],
            original_attempts=deepcopy(saved["attempts"]),
            raw_path=raw_name,
            raw_sha256=sha(raw),
            source_version="sha256:" + sha(raw),
            final_url=selected["final_url"],
            selected_sequence=selected["sequence"],
            reading={
                "status": "inaccessible",
                "evidence_level": "metadata",
                "identity_status": "unverified",
                "characters": 0,
                "text_sha256": None,
                "locators": [],
                "diagnostics": {},
                "error": {
                    "type": "PermissionError",
                    "message": "saved HTTP response was not successful; no new access attempted",
                },
            },
            source_metadata={},
            metadata={
                "journal": "Unknown",
                "volume": None,
                "issue": None,
                "pages": None,
                "doi": None,
            },
            metadata_provenance={
                "journal": {"status": "original-verbatim", "value": "Unknown"}
            },
        )
        row["attempt_id"] = sha(
            canonical(
                {
                    "base": base_sha,
                    "parser": reader.parser_sha,
                    "work": row["work_id"],
                    "version": row["version_id"],
                    "raw": row["raw_sha256"],
                }
            )
        )
        rows.append(row)
    manifest = {
        "data": {
            "kind": "Stage1SavedSourceRerun",
            "schema_version": "1.0.0",
            "base_index_sha256": base_sha,
            "parser_runtime": {
                "python": "3.11.9",
                "python_executable_sha256": "21bb438c0d4a6f1f164b9a646f6ee000340185e5871180aec06db8d3f07c0082",
                "parser_sha256": reader.parser_sha,
                "adapter_sha256": "b6a0a5eeadd70ccbbc69a0e53a1872d154447586c6ea8a67ea90ab8a847ddf7c",
                "dependencies": {
                    "pdfplumber": "0.11.7",
                    "pdfminer.six": "20250506",
                    "research-hub-pipeline": "1.2.0",
                },
            },
            "rows": rows,
            "research_execution": False,
            "scientific_judgments_changed": False,
            "official_stage2_import_eligible": False,
            "new_searches": 0,
            "new_downloads": 0,
            "quality_score": None,
        },
        "files": {
            name: {"sha256": sha(raw), "bytes": len(raw)}
            for name, raw in sorted(files.items())
        },
    }
    output.mkdir()
    for name, raw in files.items():
        path = output / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(raw)
    manifest_bytes = canonical(manifest)
    (output / "source-rerun-manifest.json").write_bytes(manifest_bytes)
    return output, sha(manifest_bytes), manifest


def _rebind_saved_import(reader, *, provenance_bytes=None, mutate_receipt=None):
    source = reader.index["sources"][0]
    receipt = source["receipt"]
    archive = reader.package / "deliverable/sources/source-0"
    mapping = artifact_map(receipt)
    provenance = archive / mapping[receipt["attempts"][0]["raw_path"]]
    if provenance_bytes is not None:
        provenance.write_bytes(provenance_bytes)
        receipt["attempts"][0]["response_bytes"] = len(provenance_bytes)
        receipt["attempts"][0]["raw_sha256"] = sha(provenance_bytes)
        receipt["diagnostics"]["saved_import"]["manifest_sha256"] = sha(
            provenance_bytes
        )
    if mutate_receipt is not None:
        mutate_receipt(receipt)
    receipt["receipt_sha256"] = receipt_digest(receipt)
    original = canonical(receipt)
    (archive / "original.json").write_bytes(original)
    source["result_sha256"] = sha(original)


def _rebind_attached_saved_import(index, *, provenance_bytes=None):
    """Rehash all outer bindings after editing one attached SDK receipt."""
    extension = index["source_rerun"]
    source = index["sources"][0]
    receipt = source["receipt"]
    row = extension["data"]["rows"][0]
    binding = row["saved_import_archive"]
    if provenance_bytes is None:
        provenance_bytes = binding["provenance_utf8"].encode("utf-8")
    receipt["attempts"][0]["response_bytes"] = len(provenance_bytes)
    receipt["attempts"][0]["raw_sha256"] = sha(provenance_bytes)
    receipt["diagnostics"]["saved_import"]["manifest_sha256"] = sha(provenance_bytes)
    receipt["receipt_sha256"] = receipt_digest(receipt)
    original = canonical(receipt)
    source["result_sha256"] = sha(original)
    row["original_result_sha256"] = source["result_sha256"]
    row["original_attempts"] = deepcopy(receipt["attempts"])
    provenance = artifact_map(receipt)[receipt["attempts"][0]["raw_path"]]
    provenance_name = binding["root"] + "/" + provenance
    original_name = binding["root"] + "/original.json"
    binding["provenance_utf8"] = provenance_bytes.decode("utf-8")
    binding["members"][provenance] = {
        "sha256": sha(provenance_bytes),
        "bytes": len(provenance_bytes),
    }
    binding["members"]["original.json"] = {
        "sha256": sha(original),
        "bytes": len(original),
    }
    extension["artifact_hashes"][provenance_name] = binding["members"][provenance]
    extension["artifact_hashes"][original_name] = binding["members"]["original.json"]
    base = deepcopy(index)
    base.pop("source_rerun")
    base["schema_version"] = (
        "2.0.0" if base["supplement"]["status"] != "not-provided" else "1.0.0"
    )
    extension["data"]["base_index_sha256"] = sha(canonical(base))
    for rerun_row in extension["data"]["rows"]:
        rerun_row["attempt_id"] = sha(
            canonical(
                {
                    "base": extension["data"]["base_index_sha256"],
                    "parser": extension["data"]["parser_runtime"]["parser_sha256"],
                    "work": rerun_row["work_id"],
                    "version": rerun_row["version_id"],
                    "raw": rerun_row.get("raw_sha256"),
                }
            )
        )
    extension["manifest_sha256"] = sha(
        canonical(
            {
                "data": extension["data"],
                "files": extension["artifact_hashes"],
            }
        )
    )
    return index


def _counterfeit_saved_import(index):
    """Rehash every outer binding around malformed author-written provenance."""
    return _rebind_attached_saved_import(deepcopy(index), provenance_bytes=b"{")


class OfflineHistoryTests(unittest.TestCase):
    def test_offline_history_replays_matching_actual_body_without_new_access(self):
        for level in ["full-text", "abstract"]:
            with self.subTest(level=level):
                offline_reader = _offline_reader(self)
                reader = offline_reader
                parser = PARSER.replace(b'"full-text"', f'"{level}"'.encode())
                reader.parser.write_bytes(parser)
                reader.parser_sha = sha(parser)
                before = deepcopy(reader.index)
                output, digest, manifest = reader.build()
                attached = attach_rerun(reader.index, output, digest)
                assert attached["source_rerun"]["data"] == manifest["data"]
                row = manifest["data"]["rows"][0]
                assert row["reading"]["status"] == "extracted"
                assert row["reading"]["evidence_level"] == level
                assert (
                    row["original_attempts"]
                    == before["sources"][0]["receipt"]["attempts"]
                )
                assert row["selected_sequence"] == 3
                assert row["acquisition_sequence"] == 2
                assert row["original_attempts"][-1]["http_status"] is None
                assert reader.index == before
                assert (
                    manifest["data"]["new_searches"]
                    == manifest["data"]["new_downloads"]
                    == 0
                )
                assert manifest["data"]["scientific_judgments_changed"] is False
                assert manifest["data"]["quality_score"] is None

    def test_offline_history_rejects_unmatched_or_unusable_acquisition(self):
        for field, value in [
            ("raw_sha256", "0" * 64),
            ("final_url", "https://example.test/other"),
            ("response_bytes", 1),
            ("response_truncated", True),
            ("outcome", "redirect"),
            ("purpose", "offline-parser-replay"),
        ]:
            with self.subTest(field=field, value=value):
                offline_reader = _offline_reader(self)
                reader = offline_reader
                reader.index["sources"][0]["receipt"]["attempts"][-2][field] = value
                _bind_history(reader)
                output, digest, manifest = reader.build()
                attach_rerun(reader.index, output, digest)
                reading = manifest["data"]["rows"][0]["reading"]
                assert reading["status"] == "failed-engineering"
                assert (
                    "no matching successful saved acquisition"
                    in reading["error"]["message"]
                )
                assert reading["characters"] == 0

    def test_offline_history_without_original_acquisition_fails_engineering(self):
        offline_reader = _offline_reader(self)
        reader = offline_reader
        reader.index["sources"][0]["receipt"]["attempts"] = [
            reader.index["sources"][0]["receipt"]["attempts"][-1]
        ]
        _bind_history(reader)
        output, digest, manifest = reader.build()
        attach_rerun(reader.index, output, digest)
        assert manifest["data"]["rows"][0]["reading"]["status"] == "failed-engineering"

    def test_offline_history_checks_actual_acquisition_archive_bytes(self):
        offline_reader = _offline_reader(self)
        reader = offline_reader
        (reader.package / "deliverable/sources/source-0/saved.html").write_bytes(
            b"altered acquisition"
        )
        output, digest, manifest = reader.build()
        attach_rerun(reader.index, output, digest)
        reading = manifest["data"]["rows"][0]["reading"]
        assert reading["status"] == "failed-engineering"
        assert "acquisition bytes differ" in reading["error"]["message"]

    def test_offline_history_missing_acquisition_archive_fails_engineering(self):
        offline_reader = _offline_reader(self)
        reader = offline_reader
        before = deepcopy(reader.index)
        archive = reader.package / "deliverable/sources/source-0"
        (archive / "saved.html").unlink()
        assert (archive / "replay.html").is_file()
        output, digest, manifest = reader.build()
        attach_rerun(reader.index, output, digest)
        row = manifest["data"]["rows"][0]
        reading = row["reading"]
        assert reading["status"] == "failed-engineering"
        assert reading["error"]["type"] == "ValueError"
        assert "acquisition archive could not be read" in reading["error"]["message"]
        assert reading["characters"] == 0 and reading["text_sha256"] is None
        assert "extracted_path" not in row
        assert row["selected_sequence"] == 3
        assert row["original_attempts"] == before["sources"][0]["receipt"]["attempts"]
        assert reader.index == before

    def test_offline_history_does_not_override_latest_actual_failure(self):
        for saved_body in [False, True]:
            with self.subTest(saved_body=saved_body):
                offline_reader = _offline_reader(self)
                reader = offline_reader
                latest = deepcopy(reader.index["sources"][0]["receipt"]["attempts"][-1])
                latest.update(
                    sequence=4,
                    purpose="synthetic actual response",
                    http_status=403 if saved_body else None,
                    outcome="http-error" if saved_body else "network-error",
                    raw_path="replay.html" if saved_body else None,
                    raw_sha256=sha(reader.raw["source-0"]) if saved_body else None,
                )
                reader.index["sources"][0]["receipt"]["attempts"].append(latest)
                _bind_history(reader)
                output, digest, manifest = reader.build()
                attach_rerun(reader.index, output, digest)
                row = manifest["data"]["rows"][0]
                assert row["reading"]["status"] == "inaccessible"
                assert row["reading"]["characters"] == 0
                assert "extracted_path" not in row
                assert row["original_attempts"][-1] == latest

    def test_offline_history_rejects_rehashed_acquisition_sequence(self):
        for sequence in [1, 3, 99]:
            with self.subTest(sequence=sequence):
                offline_reader = _offline_reader(self)
                reader = offline_reader
                output, _, manifest = reader.build()
                manifest["data"]["rows"][0]["acquisition_sequence"] = sequence
                raw = canonical(manifest)
                (output / "source-rerun-manifest.json").write_bytes(raw)
                with self.assertRaisesRegex(
                    DeliverableError, "acquisition attempt differs"
                ):
                    attach_rerun(reader.index, output, sha(raw))


class SavedPublicImportTests(unittest.TestCase):
    def test_persisted_saved_import_rejects_rehashed_identity_aliases(self):
        reader = _saved_import_reader(self)
        output, digest, _ = reader.build("identity-aliases")
        attached = attach_rerun(reader.index, output, digest)
        index_path = reader.root / "attached-index.json"
        index_path.write_bytes(canonical(attached))
        original = json.loads(index_path.read_bytes())
        self.assertEqual(validate_rerun_index(original), original)
        self.assertEqual(view_fixture.validate_index(original), original)

        for mutation in (
            "source-url",
            "expected-doi",
            "expected-title",
            "missing-source-url",
            "missing-expected-identity",
            "noncanonical-request-doi",
        ):
            mutated = deepcopy(original)
            receipt = mutated["sources"][0]["receipt"]
            if mutation == "source-url":
                receipt["source_url"] = "https://example.test/different-source"
            elif mutation == "expected-doi":
                receipt["expected_identity"]["doi"] = "10.1000/different"
            elif mutation == "expected-title":
                receipt["expected_identity"]["title"] = "Different evidence"
            elif mutation == "missing-source-url":
                receipt.pop("source_url")
            elif mutation == "missing-expected-identity":
                receipt.pop("expected_identity")
            else:
                receipt["request"]["doi"] = "https://doi.org/10.1000/synthetic"
            _rebind_attached_saved_import(mutated)
            mutated_path = reader.root / f"{mutation}-index.json"
            mutated_path.write_bytes(canonical(mutated))
            persisted = json.loads(mutated_path.read_bytes())
            for name, gate in (
                ("rerun-index", validate_rerun_index),
                ("workspace-view", view_fixture.validate_index),
            ):
                with self.subTest(mutation=mutation, gate=name):
                    with self.assertRaisesRegex(
                        DeliverableError,
                        "saved import provenance/receipt binding differs",
                    ):
                        gate(persisted)

    def test_saved_import_without_claims_exports_without_fabricated_rows(self):
        reader = _saved_import_reader(self)
        self.assertEqual(reader.index["claims"], [])
        output, digest, _ = reader.build("empty-claims-export")
        attached = attach_rerun(reader.index, output, digest)
        exported = rerun_files(attached, output)
        availability = json.loads(exported["source-rerun/source-availability.json"])
        self.assertEqual(availability["claim_rows"], [])
        workbook = load_workbook(BytesIO(exported["source-rerun/catalog.xlsx"]))
        self.addCleanup(workbook.close)
        self.assertIn("SourceRerun", workbook.sheetnames)
        self.assertIn("SourceAvailability", workbook.sheetnames)
        self.assertNotIn("ClaimAvailability", workbook.sheetnames)

    def test_malformed_saved_import_diagnostics_are_rejected(self):
        reader = _saved_import_reader(self)
        original = reader.index["sources"][0]["receipt"]
        for diagnostics in [None, [], {"saved_import": None}, {"saved_import": []}]:
            with self.subTest(diagnostics=diagnostics):
                receipt = deepcopy(original)
                receipt["diagnostics"] = diagnostics
                self.assertFalse(_saved_public_import(receipt, receipt["attempts"][1]))

    def test_saved_import_rejects_incomplete_archive_binding(self):
        reader = _saved_import_reader(self)
        output, _, manifest = _legacy_saved_import_rerun(reader)
        manifest["data"]["rows"][0]["saved_import_archive"] = {}
        raw = canonical(manifest)
        (output / "source-rerun-manifest.json").write_bytes(raw)
        with self.assertRaisesRegex(
            DeliverableError, "saved import archive root differs"
        ):
            attach_rerun(reader.index, output, sha(raw))

    def test_complete_legacy_saved_import_stays_inaccessible_in_view_and_export(self):
        reader = _saved_import_reader(self)
        reader.index["claims"] = [
            {
                "claim_id": "historical-claim",
                "work_id": "fixture-0",
                "version_id": "version-0",
                "text": "A historical synthetic claim remains unmodified.",
                "support": "Unknown",
            }
        ]
        reader.index["coverage"] = [
            {"need_id": "synthetic-need", "status": "unresolved", "coverage": "Unknown"}
        ]
        before = deepcopy(reader.index)
        output, digest, manifest = _legacy_saved_import_rerun(reader)
        legacy = deepcopy(manifest)
        attached = attach_rerun(reader.index, output, digest)
        self.assertEqual(view_fixture.validate_index(attached), attached)
        self.assertEqual(attached["source_rerun"]["data"], legacy["data"])
        self.assertEqual(reader.index, before)
        self.assertEqual(
            {key: attached[key] for key in before if key != "schema_version"},
            {key: before[key] for key in before if key != "schema_version"},
        )
        self.assertEqual(attached["status"], "partial-review-only")
        self.assertEqual(
            [paper["venue"] for paper in attached["papers"]], ["Unknown"] * 2
        )
        self.assertEqual(attached["claims"][0]["support"], "Unknown")
        self.assertEqual(attached["coverage"][0]["coverage"], "Unknown")
        self.assertEqual(attached["coverage"][0]["status"], "unresolved")
        for row, original in zip(legacy["data"]["rows"], before["sources"]):
            self.assertNotIn("saved_import_archive", row)
            self.assertNotIn("extracted_path", row)
            self.assertEqual(row["original_attempts"], original["receipt"]["attempts"])
            self.assertEqual(row["reading"]["status"], "inaccessible")
            self.assertEqual(row["reading"]["characters"], 0)
        exported = rerun_files(attached, output)
        self.assertEqual(
            json.loads(exported["source-rerun/manifest.json"]), attached["source_rerun"]
        )
        self.assertFalse(any(name.endswith("extracted.txt") for name in exported))
        index_bytes = canonical(attached)
        index_path = reader.root / "legacy-index.json"
        index_path.write_bytes(index_bytes)
        view = reader.root / "legacy-view"
        reference = Path(__file__).parents[1] / "references/research-workspace"
        view_fixture.render_view(index_path, view, reference, sha(index_bytes))
        self.assertEqual((view / "workspace-index.json").read_bytes(), index_bytes)
        self.assertEqual(
            (view / "source-rerun/manifest.json").read_bytes(),
            exported["source-rerun/manifest.json"],
        )
        self.assertEqual(manifest, legacy)

    def test_rehashed_legacy_saved_import_cannot_promote_text_without_archive(self):
        reader = _saved_import_reader(self)
        output, _, manifest = _legacy_saved_import_rerun(reader)
        row = manifest["data"]["rows"][0]
        text = b"Counterfeit admitted text"
        row["extracted_path"] = "sources/source-0/extracted.txt"
        row["reading"].update(
            status="extracted",
            evidence_level="full-text",
            identity_status="consistent",
            characters=len(text),
            text_sha256=sha(text),
            error=None,
        )
        (output / row["extracted_path"]).write_bytes(text)
        manifest["files"][row["extracted_path"]] = {
            "sha256": sha(text),
            "bytes": len(text),
        }
        raw = canonical(manifest)
        (output / "source-rerun-manifest.json").write_bytes(raw)
        with self.assertRaisesRegex(DeliverableError, "unbound saved import"):
            attach_rerun(reader.index, output, sha(raw))

    def test_successful_saved_import_cannot_strip_archive_binding(self):
        reader = _saved_import_reader(self)
        output, _, manifest = reader.build("stripped-archive-binding")
        row = manifest["data"]["rows"][0]
        self.assertEqual(row["reading"]["status"], "extracted")
        row.pop("saved_import_archive")
        raw = canonical(manifest)
        (output / "source-rerun-manifest.json").write_bytes(raw)
        with self.assertRaisesRegex(DeliverableError, "unbound saved import"):
            attach_rerun(reader.index, output, sha(raw))

    def test_legacy_saved_import_rejects_rehashed_raw_substitution_in_all_gates(self):
        for mutation in (
            "missing-row-hash",
            "empty-row-hash",
            "substituted-raw",
            "wrong-byte-count",
        ):
            with self.subTest(mutation=mutation):
                reader = _saved_import_reader(self)
                output, digest, _ = _legacy_saved_import_rerun(reader)
                attached = attach_rerun(reader.index, output, digest)
                original = deepcopy(attached["sources"][0])
                extension = attached["source_rerun"]
                row = extension["data"]["rows"][0]
                raw_path = output / row["raw_path"]
                if mutation == "missing-row-hash":
                    row.pop("raw_sha256")
                elif mutation == "empty-row-hash":
                    row["raw_sha256"] = ""
                elif mutation == "substituted-raw":
                    replacement = b"X" * len(raw_path.read_bytes())
                    raw_path.write_bytes(replacement)
                    extension["artifact_hashes"][row["raw_path"]]["sha256"] = sha(
                        replacement
                    )
                else:
                    extension["artifact_hashes"][row["raw_path"]]["bytes"] += 1
                row["attempt_id"] = sha(
                    canonical(
                        {
                            "base": extension["data"]["base_index_sha256"],
                            "parser": extension["data"]["parser_runtime"][
                                "parser_sha256"
                            ],
                            "work": row["work_id"],
                            "version": row["version_id"],
                            "raw": row.get("raw_sha256"),
                        }
                    )
                )
                extension["manifest_sha256"] = sha(
                    canonical(
                        {
                            "data": extension["data"],
                            "files": extension["artifact_hashes"],
                        }
                    )
                )
                self.assertEqual(attached["sources"][0], original)
                self.assertEqual(row["reading"]["status"], "inaccessible")
                self.assertEqual(row["reading"]["characters"], 0)
                self.assertEqual(
                    row["source_version"], original["receipt"]["source_version"]
                )
                with self.assertRaisesRegex(
                    DeliverableError, "legacy saved import raw binding differs"
                ):
                    validate_rerun_index(attached)
                with self.assertRaisesRegex(
                    DeliverableError, "legacy saved import raw binding differs"
                ):
                    view_fixture.validate_index(attached)
                with self.assertRaisesRegex(
                    DeliverableError, "legacy saved import raw binding differs"
                ):
                    rerun_files(attached, output)
                index_bytes = canonical(attached)
                index_path = reader.root / "rehashed-index.json"
                index_path.write_bytes(index_bytes)
                reference = Path(__file__).parents[1] / "references/research-workspace"
                with self.assertRaisesRegex(
                    DeliverableError, "legacy saved import raw binding differs"
                ):
                    view_fixture.render_view(
                        index_path,
                        reader.root / "rejected-view",
                        reference,
                        sha(index_bytes),
                    )

    def test_saved_import_rejects_rehashed_separate_raw_in_all_gates(self):
        reader = _saved_import_reader(self)
        reader.index["claims"] = [
            {
                "claim_id": "saved-source-claim",
                "work_id": "fixture-0",
                "version_id": "version-0",
                "text": "A synthetic historical claim keeps its source binding.",
                "support": "Unknown",
            }
        ]
        output, digest, _ = reader.build("separate-raw-substitution")
        attached = attach_rerun(reader.index, output, digest)
        extension = attached["source_rerun"]
        row = extension["data"]["rows"][0]
        raw_path = output / row["raw_path"]
        replacement = b"X" * len(raw_path.read_bytes())
        self.assertNotEqual(sha(replacement), row["raw_sha256"])
        raw_path.write_bytes(replacement)
        extension["artifact_hashes"][row["raw_path"]]["sha256"] = sha(replacement)
        manifest_bytes = canonical(
            {"data": extension["data"], "files": extension["artifact_hashes"]}
        )
        extension["manifest_sha256"] = sha(manifest_bytes)
        (output / "source-rerun-manifest.json").write_bytes(manifest_bytes)
        for gate in (validate_rerun_index, view_fixture.validate_index):
            with self.subTest(gate=gate.__module__):
                with self.assertRaises(DeliverableError):
                    gate(attached)
        with self.subTest(gate="export"):
            with self.assertRaises(DeliverableError):
                rerun_files(attached, output)
        with self.subTest(gate="attachment"):
            with self.assertRaises(DeliverableError):
                attach_rerun(reader.index, output, sha(manifest_bytes))
        index_bytes = canonical(attached)
        index_path = reader.root / "substituted-index.json"
        index_path.write_bytes(index_bytes)
        reference = Path(__file__).parents[1] / "references/research-workspace"
        with self.subTest(gate="render"):
            with self.assertRaises(DeliverableError):
                view_fixture.render_view(
                    index_path,
                    reader.root / "rejected-view",
                    reference,
                    sha(index_bytes),
                )

    def test_saved_public_import_replays_sdk_verified_source_without_fake_http(self):
        reader = _saved_import_reader(self)
        legacy_sources = reader.package / "deliverable/sources"
        legacy_sources.rename(reader.package / "sources")
        before = deepcopy(reader.index)
        output, digest, manifest = reader.build()
        attach_rerun(reader.index, output, digest)
        row = manifest["data"]["rows"][0]
        assert row["reading"]["status"] == "extracted"
        assert row["selected_sequence"] == 2
        assert row["original_attempts"] == before["sources"][0]["receipt"]["attempts"]
        assert row["original_attempts"][0]["purpose"] == "saved-input-provenance"
        assert row["original_attempts"][1]["purpose"] == "saved-public-source"
        assert row["original_attempts"][1]["http_status"] is None
        assert "acquisition_sequence" not in row
        archive = row["saved_import_archive"]
        assert set(archive["members"]) == {
            "original.json",
            *artifact_map(reader.index["sources"][0]["receipt"]).values(),
        }
        assert reader.index == before

    def test_saved_public_import_in_memory_and_export_gates_fail_closed(self):
        reader = _saved_import_reader(self)
        output, digest, _ = reader.build("counterfeit-index")
        attached = attach_rerun(reader.index, output, digest)
        counterfeit = _counterfeit_saved_import(attached)
        with self.assertRaisesRegex(
            DeliverableError, "saved import provenance is invalid"
        ):
            validate_rerun_index(counterfeit)
        with self.assertRaisesRegex(
            DeliverableError, "saved import provenance is invalid"
        ):
            rerun_files(counterfeit, output)

        receipt = attached["sources"][0]["receipt"]
        row = attached["source_rerun"]["data"]["rows"][0]
        selected = artifact_map(receipt)[receipt["attempts"][1]["raw_path"]]
        selected_path = output / row["saved_import_archive"]["root"] / selected
        selected_path.write_bytes(b"tampered contained source")
        with self.assertRaisesRegex(DeliverableError, "bytes changed while exporting"):
            rerun_files(attached, output)

    def test_saved_public_import_attach_replays_contained_sdk_archive(self):
        def rewrite(output, manifest):
            raw = canonical(manifest)
            (output / "source-rerun-manifest.json").write_bytes(raw)
            return sha(raw)

        reader = _saved_import_reader(self)
        output, _, manifest = reader.build("attach-malformed-provenance")
        row = manifest["data"]["rows"][0]
        receipt = reader.index["sources"][0]["receipt"]
        provenance = artifact_map(receipt)[receipt["attempts"][0]["raw_path"]]
        name = row["saved_import_archive"]["root"] + "/" + provenance
        malformed = b"{"
        (output / name).write_bytes(malformed)
        bound = {"sha256": sha(malformed), "bytes": len(malformed)}
        manifest["files"][name] = bound
        row["saved_import_archive"]["members"][provenance] = bound
        with self.assertRaisesRegex(
            DeliverableError, "saved import provenance byte binding differs"
        ):
            attach_rerun(reader.index, output, rewrite(output, manifest))

        reader = _saved_import_reader(self)
        output, _, manifest = reader.build("attach-missing-provenance")
        row = manifest["data"]["rows"][0]
        receipt = reader.index["sources"][0]["receipt"]
        provenance = artifact_map(receipt)[receipt["attempts"][0]["raw_path"]]
        name = row["saved_import_archive"]["root"] + "/" + provenance
        (output / name).unlink()
        manifest["files"].pop(name)
        row["saved_import_archive"]["members"].pop(provenance)
        with self.assertRaisesRegex(
            DeliverableError, "saved import archive inventory differs"
        ):
            attach_rerun(reader.index, output, rewrite(output, manifest))

    def test_saved_public_import_provenance_and_attempt_corruption_fail_closed(self):
        def changed_manifest(reader, **updates):
            receipt = reader.index["sources"][0]["receipt"]
            mapping = artifact_map(receipt)
            path = (
                reader.package
                / "deliverable/sources/source-0"
                / mapping[receipt["attempts"][0]["raw_path"]]
            )
            value = json.loads(path.read_bytes())
            value.update(updates)
            return canonical(value)

        cases = {
            "malformed-provenance": lambda reader: _rebind_saved_import(
                reader, provenance_bytes=b"{"
            ),
            "wrong-source-hash": lambda reader: _rebind_saved_import(
                reader,
                provenance_bytes=changed_manifest(reader, raw_sha256="0" * 64),
            ),
            "wrong-manifest-version": lambda reader: _rebind_saved_import(
                reader,
                provenance_bytes=changed_manifest(
                    reader, schema_version="saved-public-source-input/v2"
                ),
            ),
            "truncated-source": lambda reader: _rebind_saved_import(
                reader,
                mutate_receipt=lambda receipt: receipt["attempts"][1].update(
                    response_truncated=True
                ),
            ),
            "sequence-offset": lambda reader: _rebind_saved_import(
                reader,
                mutate_receipt=lambda receipt: receipt["attempts"][1].update(
                    sequence=3
                ),
            ),
        }
        for name, mutate in cases.items():
            with self.subTest(case=name):
                reader = _saved_import_reader(self)
                mutate(reader)
                with self.assertRaisesRegex(
                    DeliverableError, "public source semantic replay failed"
                ):
                    reader.build(name)


if __name__ == "__main__":
    unittest.main()
