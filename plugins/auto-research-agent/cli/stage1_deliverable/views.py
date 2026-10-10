"""Deterministic editable views; these never become canonical input."""

import csv
import html
import io
import json
import zipfile
from datetime import datetime
from urllib.parse import quote

from docx import Document
from docx.shared import RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .common import DeliverableError, canonical


def cell(value):
    if value is None:
        return "unknown"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _zip_bytes(raw):
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return out.getvalue()


def tables(records, sources):
    papers, classification, findings = [], [], []
    for row in records["papers"]:
        attempts = [s for s in sources if s["work_id"] == row["work_id"]]
        identity = {
            k: row[k]
            for k in (
                "work_id",
                "version_id",
                "title",
                "authors",
                "year",
                "venue",
                "doi",
                "url",
                "evidence_level",
            )
        }
        papers.append(
            {
                **identity,
                "access_dates": [x["observed_at"] for x in attempts],
                "source_ids": row["source_ids"],
            }
        )
        key = {k: row[k] for k in ("work_id", "version_id")}
        classification.append({**key, **row["classification"], "roles": row["roles"]})
        findings.append({**key, **row["findings"], "claim_ids": row["claim_ids"]})
    claims = []
    for row in records["claims"]:
        source = next(
            x
            for x in sources
            if x["source_id"] == row["source_id"] and x["source_version"]
        )
        claims.append(
            {
                **row,
                "source_sha256": source["sha256"],
                "source_version": source["source_version"],
            }
        )
    return {
        "Papers": papers,
        "Classification": classification,
        "Findings": findings,
        "Claims": claims,
        "Screening": records["screening"],
        "Coverage": records["coverage"],
    }


EMPTY_HEADERS = {
    "Claims": [
        "claim_id",
        "work_id",
        "version_id",
        "text",
        "source_id",
        "relation",
        "evidence_level",
        "locator",
        "start",
        "end",
        "quote",
        "source_sha256",
        "source_version",
    ]
}


def grid(name, rows):
    headers = list(rows[0]) if rows else EMPTY_HEADERS[name]
    return [headers, *[[cell(row[key]) for key in headers] for row in rows]]


def csv_bytes(name, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    for row in grid(name, rows):
        writer.writerow(
            [
                "'" + value
                if value.lstrip().startswith(("=", "+", "-", "@"))
                else value
                for value in row
            ]
        )
    return stream.getvalue().encode("utf-8-sig")


def workbook_bytes(all_tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.created = workbook.properties.modified = datetime(2000, 1, 1)
    for name, rows in all_tables.items():
        sheet = workbook.create_sheet(name)
        for row in grid(name, rows):
            if any(len(value.encode("utf-16-le")) // 2 > 32767 for value in row):
                raise DeliverableError("canonical value exceeds Excel cell limit")
            sheet.append(row)
        for row in sheet:
            for item in row:
                item.data_type = "s"  # Prevent external-data formula execution.
                item.alignment = Alignment(vertical="top", wrap_text=True)
        for item in sheet[1]:
            item.font = Font(bold=True, color="FFFFFF")
            item.fill = PatternFill("solid", fgColor="204A63")
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for number in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(number)].width = (
                24 if number <= 2 else 44
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    # openpyxl updates modified at save; remove wall-clock entropy explicitly.
    source = io.BytesIO(buffer.getvalue())
    out = io.BytesIO()
    with zipfile.ZipFile(source) as z, zipfile.ZipFile(out, "w") as target:
        for name in z.namelist():
            raw = z.read(name)
            if name == "docProps/core.xml":
                import re

                raw = re.sub(
                    rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\2",
                    raw,
                )
            target.writestr(name, raw)
    return _zip_bytes(out.getvalue())


def review_lines(records):
    lines = [
        "# Literature review",
        "",
        records["topic"],
        "",
        "As of: " + records["as_of"],
        "",
        "Experimental researcher package. Source and interpretation checks remain separate; no P1-P3 improvement is established.",
        "",
    ]
    for paper in records["papers"]:
        lines.extend(
            [
                "## " + paper["work_id"] + " — " + paper["title"],
                "",
                "Version: " + paper["version_id"],
                "Evidence level: " + paper["evidence_level"],
                "",
            ]
        )
        for key, value in paper["findings"].items():
            lines.extend([key.replace("_", " ").capitalize() + ": " + value, ""])
        lines.extend(
            [
                "Claims: " + ", ".join(paper["claim_ids"]),
                "Roles (authored judgments, not exporter endorsements): "
                + cell(paper["roles"]),
                "",
            ]
        )
    lines.extend(["## Evidence and limitations", ""])
    for claim in records["claims"]:
        lines.extend(
            [
                f"{claim['claim_id']} [{claim['work_id']} / {claim['version_id']}]: {claim['text']}",
                f"{claim['relation']}; {claim['evidence_level']}; {claim['locator']}; {claim['source_id']} characters {claim['start']}:{claim['end']}",
                "Exact source excerpt retained unchanged: workbook Claims sheet "
                f"and canonical records, claim_id {claim['claim_id']}. "
                "CSV cells use formula-safe escaping. Original extraction "
                "spacing and character offsets are preserved for audit.",
                "",
            ]
        )
    return lines


def docx_bytes(lines):
    document = Document()
    document.core_properties.created = document.core_properties.modified = datetime(
        2000, 1, 1
    )
    document.core_properties.title = "Stage 1 literature review"
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    for line in lines:
        if line.startswith("# "):
            paragraph = document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            paragraph = document.add_heading(line[3:], 1)
        else:
            document.add_paragraph(line)
            continue
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    for element in (document.element, document.styles.element):
        for border in list(element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    output = io.BytesIO()
    document.save(output)
    return _zip_bytes(output.getvalue())


def bibtex(records):
    def escape(value):
        # Translate once: generated TeX syntax must never be escaped again.
        # Named brace glyphs keep even an unmatched input brace BibTeX-safe.
        replacements = {
            "\\": r"{\textbackslash}",
            "{": r"{\textbraceleft}",
            "}": r"{\textbraceright}",
            "%": r"\%",
            "&": r"\&",
            "_": r"\_",
            "#": r"\#",
            "$": r"\$",
            "~": r"{\textasciitilde}",
            "^": r"{\textasciicircum}",
            "\n": " ",
            "\r": " ",
            "\t": " ",
        }
        return "".join(replacements.get(char, char) for char in cell(value))

    entries = []
    for paper in records["papers"]:
        fields = {
            "title": paper["title"],
            "author": " and ".join(paper["authors"]),
            "year": paper["year"],
            "journal": paper["venue"],
            "doi": paper["doi"],
            "url": paper["url"],
            "note": "Version "
            + paper["version_id"]
            + "; evidence "
            + paper["evidence_level"],
        }
        lines = ["@misc{" + paper["work_id"] + ","]
        for key, value in fields.items():
            if value is None:
                continue
            # BibLaTeX URL/DOI fields are verbatim, not TeX text. Preserve URI
            # separators and existing percent escapes; encode unsafe braces.
            encoded = (
                quote(value, safe=":/?#[]@!$&'()*+,;=%-._~")
                if key in {"url", "doi"}
                else escape(value)
            )
            lines.append(f"  {key} = {{{encoded}}},")
        entries.append("\n".join([*lines, "}"]))
    return ("\n\n".join(entries) + "\n").encode("utf-8")


def render(records, sources):
    # JSON object order is not research data. Match the manifest's canonical
    # ordering while preserving authored array order and original input bytes.
    records = json.loads(canonical(records))
    all_tables = tables(records, sources)
    lines = review_lines(records)
    markdown = "\n".join(html.escape(line, quote=False) for line in lines) + "\n"
    coverage = [
        "# Coverage and stopping record",
        "",
        "Recorded judgments; no automated scientific sufficiency claim.",
        "",
    ]
    for row in records["coverage"]:
        coverage.extend(
            [
                "## " + row["need_id"],
                *[
                    key + ": " + html.escape(cell(value), quote=False)
                    for key, value in row.items()
                ],
                "",
            ]
        )
    return {
        "literature_catalog.xlsx": workbook_bytes(all_tables),
        "literature_review.md": markdown.encode("utf-8"),
        "literature_review.docx": docx_bytes(lines),
        "references.bib": bibtex(records),
        "claims_and_evidence.csv": csv_bytes("Claims", all_tables["Claims"]),
        "search_and_screening.csv": csv_bytes("Screening", all_tables["Screening"]),
        "coverage_and_stop.md": ("\n".join(coverage) + "\n").encode("utf-8"),
        "README.md": b"# Stage 1 researcher deliverable\n\nExperimental; quality improvement not yet demonstrated.\n\npapers.jsonl and provenance_manifest.json contain canonical records. Excel, Markdown, DOCX, BibTeX and CSV are editable derived views. Editing a view does not update canonical research records; regeneration and validation are required. Source archives preserve all attempts, including unavailable evidence. Papers contain only selected public full sources. Do not commit this private package to Git.\n\nWord and Markdown show authored findings, claim statements and source locators. Find each exact excerpt by claim_id in the workbook Claims sheet and canonical records. claims_and_evidence.csv protects formula-like values with a leading apostrophe, so its displayed quote may differ. Original quotes, extraction spacing and character offsets remain bound to the source; presentation does not establish claim support.\n\nValidate with: python -m stage1_deliverable validate PACKAGE --expected-sha256 MANIFEST_SHA256\nThe trusted manifest hash must come from the separately retained export receipt. Validation proves bindings and reproducible views, not scientific correctness. DOCX is supported by the declared runtime.\n",
    }
